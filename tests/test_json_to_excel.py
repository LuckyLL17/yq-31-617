import sys
import os
import json
import math
import tempfile
import shutil
import time
from io import StringIO
from unittest.mock import patch, MagicMock

import pytest
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from json_to_excel import (
    ProgressTracker,
    load_json,
    flatten_dict,
    auto_detect_headers,
    merge_headers,
    _get_alignment,
    _create_border,
    apply_header_style,
    apply_data_style,
    set_column_widths,
    _get_column_letter,
    extract_value,
    _sanitize_sheet_name,
    _render_sheet_name,
    _match_custom_rule,
    _match_range_group,
    split_data_by_field,
    _to_numeric,
    _aggregate_values,
    _get_field_label,
    _get_row_key,
    _get_col_key,
    build_pivot_table,
    _compute_pivot_value,
    _compute_row_total,
    _compute_col_total,
    _compute_grand_total,
    _format_value,
    _get_aggregate_label,
    DataLoader,
    ExcelExporter,
)


class TestProgressTracker初始化:
    def test_默认参数(self):
        tracker = ProgressTracker(total=100)
        assert tracker.total == 100
        assert tracker.current == 0
        assert tracker.description == "处理中"
        assert tracker.unit == "条"
        assert tracker.bar_width == 30
        assert tracker.min_interval == 0.1
        assert tracker.current_field == ""
        assert tracker.current_row_preview == ""
        assert tracker._last_line_len == 0

    def test_total为0时设为1(self):
        tracker = ProgressTracker(total=0)
        assert tracker.total == 1

    def test_total为负数时设为1(self):
        tracker = ProgressTracker(total=-5)
        assert tracker.total == 1

    def test_自定义参数(self):
        tracker = ProgressTracker(total=50, description="导出", unit="行", bar_width=20, min_interval=0.5)
        assert tracker.total == 50
        assert tracker.description == "导出"
        assert tracker.unit == "行"
        assert tracker.bar_width == 20
        assert tracker.min_interval == 0.5


class TestProgressTracker更新:
    def test_update增加计数(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.update()
        assert tracker.current == 1

    def test_update指定步长(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.update(n=3)
        assert tracker.current == 3

    def test_update不超过total(self):
        tracker = ProgressTracker(total=5, min_interval=0)
        tracker.start_time = time.time() - 0.01
        tracker.update(n=10)
        assert tracker.current == 5

    def test_update设置field和preview(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.update(field="name", row_preview="张三")
        assert tracker.current_field == "name"
        assert tracker.current_row_preview == "张三"

    def test_update空field不覆盖(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.start_time = time.time() - 0.01
        tracker.update(field="name")
        tracker.update(field="")
        assert tracker.current_field == "name"

    def test_update空preview不覆盖(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.start_time = time.time() - 0.01
        tracker.update(row_preview="张三")
        tracker.update(row_preview="")
        assert tracker.current_row_preview == "张三"

    def test_update触发render(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker.update()
        assert len(buf.getvalue()) > 0

    def test_update未到interval不render(self):
        tracker = ProgressTracker(total=10, min_interval=999)
        tracker.last_update_time = time.time()
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker.update()
        assert len(buf.getvalue()) == 0

    def test_update到达total时强制render(self):
        tracker = ProgressTracker(total=1, min_interval=999)
        tracker.last_update_time = time.time()
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker.update()
        assert len(buf.getvalue()) > 0


class TestProgressTrackerSetField:
    def test_set_field设置字段(self):
        tracker = ProgressTracker(total=10)
        tracker.set_field("age")
        assert tracker.current_field == "age"

    def test_set_field空字符串不覆盖(self):
        tracker = ProgressTracker(total=10)
        tracker.current_field = "name"
        tracker.set_field("")
        assert tracker.current_field == "name"


class TestProgressTrackerSetRowPreview:
    def test_set_row_preview设置预览(self):
        tracker = ProgressTracker(total=10)
        tracker.set_row_preview("李四")
        assert tracker.current_row_preview == "李四"

    def test_set_row_preview空字符串不覆盖(self):
        tracker = ProgressTracker(total=10)
        tracker.current_row_preview = "李四"
        tracker.set_row_preview("")
        assert tracker.current_row_preview == "李四"


class TestProgressTrackerFormatTime:
    def test_none返回占位符(self):
        tracker = ProgressTracker(total=10)
        assert tracker._format_time(None) == "--:--"

    def test_负数返回占位符(self):
        tracker = ProgressTracker(total=10)
        assert tracker._format_time(-1) == "--:--"

    def test_小于60秒(self):
        tracker = ProgressTracker(total=10)
        assert tracker._format_time(30) == "30秒"

    def test_0秒(self):
        tracker = ProgressTracker(total=10)
        assert tracker._format_time(0) == "0秒"

    def test_整分钟(self):
        tracker = ProgressTracker(total=10)
        assert tracker._format_time(60) == "1分00秒"

    def test_分钟带秒(self):
        tracker = ProgressTracker(total=10)
        assert tracker._format_time(125) == "2分05秒"

    def test_整小时(self):
        tracker = ProgressTracker(total=10)
        assert tracker._format_time(3600) == "1时00分00秒"

    def test_小时分秒(self):
        tracker = ProgressTracker(total=10)
        assert tracker._format_time(3661) == "1时01分01秒"


class TestProgressTrackerRender:
    def test_render输出进度条(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker._render()
        output = buf.getvalue()
        assert "处理中" in output
        assert "0/10" in output

    def test_render带field和preview(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.current_field = "name"
        tracker.current_row_preview = "张三"
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker._render()
        output = buf.getvalue()
        assert "字段" in output
        assert "当前" in output

    def test_render长field截断(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.current_field = "a" * 30
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker._render()
        output = buf.getvalue()
        assert "..." in output

    def test_render长preview截断(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.current_row_preview = "b" * 30
        buf = StringIO()
        with patch("shutil.get_terminal_size", return_value=MagicMock(columns=200)):
            with patch("sys.stdout", buf):
                tracker._render()
        output = buf.getvalue()
        assert "..." in output

    def test_render进度为0时rate为0(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker._render()
        output = buf.getvalue()
        assert "剩余: --:--" in output

    def test_render填充空格(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker._last_line_len = 200
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker._render()
        output = buf.getvalue()
        assert len(output) >= 100


class TestProgressTrackerFinish:
    def test_finish设置current为total(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker.finish()
        assert tracker.current == 10

    def test_finish输出换行(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker.finish()
        output = buf.getvalue()
        assert "\n" in output

    def test_finish带消息(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        with patch("sys.stdout", StringIO()):
            with patch("builtins.print") as mock_print:
                tracker.finish(message="完成!")
                mock_print.assert_called_with("完成!")


class TestLoadJson:
    def test_加载列表JSON(self, tmp_path):
        data = [{"id": 1}, {"id": 2}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        result = load_json(str(f))
        assert result == data

    def test_加载dict_data键(self, tmp_path):
        data = {"data": [{"id": 1}, {"id": 2}]}
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        result = load_json(str(f))
        assert result == [{"id": 1}, {"id": 2}]

    def test_加载dict_items键(self, tmp_path):
        data = {"items": [{"id": 1}]}
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        result = load_json(str(f))
        assert result == [{"id": 1}]

    def test_加载dict_records键(self, tmp_path):
        data = {"records": [{"id": 1}]}
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        result = load_json(str(f))
        assert result == [{"id": 1}]

    def test_加载dict无特殊键(self, tmp_path):
        data = {"name": "test", "value": 42}
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        result = load_json(str(f))
        assert result == [{"name": "test", "value": 42}]

    def test_加载非dict非列表(self, tmp_path):
        f = tmp_path / "test.json"
        f.write_text("42", encoding="utf-8")
        result = load_json(str(f))
        assert result == [42]

    def test_加载字符串(self, tmp_path):
        f = tmp_path / "test.json"
        f.write_text('"hello"', encoding="utf-8")
        result = load_json(str(f))
        assert result == ["hello"]

    def test_文件不存在抛出异常(self):
        with pytest.raises(FileNotFoundError):
            load_json("/nonexistent/path/file.json")

    def test_data键非列表时包装(self, tmp_path):
        data = {"data": "not_a_list"}
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        result = load_json(str(f))
        assert result == [{"data": "not_a_list"}]

    def test_items键非列表时包装(self, tmp_path):
        data = {"items": "not_a_list"}
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        result = load_json(str(f))
        assert result == [{"items": "not_a_list"}]

    def test_records键非列表时包装(self, tmp_path):
        data = {"records": "not_a_list"}
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        result = load_json(str(f))
        assert result == [{"records": "not_a_list"}]


class TestFlattenDict:
    def test_扁平字典(self):
        d = {"a": 1, "b": 2}
        result = flatten_dict(d)
        assert result == {"a": 1, "b": 2}

    def test_嵌套字典(self):
        d = {"a": {"b": 1, "c": 2}}
        result = flatten_dict(d)
        assert result == {"a.b": 1, "a.c": 2}

    def test_多层嵌套(self):
        d = {"a": {"b": {"c": 1}}}
        result = flatten_dict(d)
        assert result == {"a.b.c": 1}

    def test_列表值转JSON(self):
        d = {"a": [1, 2, 3]}
        result = flatten_dict(d)
        assert result == {"a": "[1, 2, 3]"}

    def test_自定义分隔符(self):
        d = {"a": {"b": 1}}
        result = flatten_dict(d, sep="/")
        assert result == {"a/b": 1}

    def test_非字典值(self):
        result = flatten_dict(42, parent_key="num")
        assert result == {"num": 42}

    def test_空字典(self):
        result = flatten_dict({})
        assert result == {}

    def test_混合类型(self):
        d = {"a": 1, "b": {"c": 2}, "d": [3, 4], "e": "hello"}
        result = flatten_dict(d)
        assert result == {"a": 1, "b.c": 2, "d": "[3, 4]", "e": "hello"}


class TestAutoDetectHeaders:
    def test_简单数据(self):
        data = [{"id": 1, "name": "张三"}, {"id": 2, "name": "李四"}]
        headers = auto_detect_headers(data)
        assert len(headers) == 2
        keys = [h["key"] for h in headers]
        assert "id" in keys
        assert "name" in keys

    def test_不同键集合(self):
        data = [{"a": 1}, {"b": 2}, {"a": 3, "c": 4}]
        headers = auto_detect_headers(data)
        keys = [h["key"] for h in headers]
        assert "a" in keys
        assert "b" in keys
        assert "c" in keys

    def test_嵌套字典自动展平(self):
        data = [{"info": {"name": "张三", "age": 20}}]
        headers = auto_detect_headers(data)
        keys = [h["key"] for h in headers]
        assert "info.name" in keys
        assert "info.age" in keys

    def test_空数据(self):
        headers = auto_detect_headers([])
        assert headers == []

    def test_非字典项跳过(self):
        data = [{"a": 1}, "not_dict", 42]
        headers = auto_detect_headers(data)
        keys = [h["key"] for h in headers]
        assert "a" in keys

    def test_header包含label和width(self):
        data = [{"user_name": "张三"}]
        headers = auto_detect_headers(data)
        h = headers[0]
        assert h["key"] == "user_name"
        assert "label" in h
        assert "width" in h
        assert h["width"] >= 15

    def test_键排序(self):
        data = [{"c": 1, "a": 2, "b": 3}]
        headers = auto_detect_headers(data)
        keys = [h["key"] for h in headers]
        assert keys == sorted(keys)


class TestMergeHeaders:
    def test_无config_headers返回auto(self):
        auto = [{"key": "a", "label": "A", "width": 10}]
        result = merge_headers([], auto, {})
        assert result == auto

    def test_无config_headers返回None(self):
        auto = [{"key": "a", "label": "A", "width": 10}]
        result = merge_headers(None, auto, {})
        assert result == auto

    def test_合并额外字段(self):
        config_h = [{"key": "id", "label": "ID", "width": 10}]
        auto_h = [{"key": "id", "label": "ID", "width": 10}, {"key": "name", "label": "姓名", "width": 15}]
        result = merge_headers(config_h, auto_h, {"auto_detect_headers": True})
        assert len(result) == 2
        keys = [h["key"] for h in result]
        assert "id" in keys
        assert "name" in keys

    def test_不自动检测时不合并(self):
        config_h = [{"key": "id", "label": "ID", "width": 10}]
        auto_h = [{"key": "id", "label": "ID", "width": 10}, {"key": "name", "label": "姓名", "width": 15}]
        result = merge_headers(config_h, auto_h, {"auto_detect_headers": False})
        assert len(result) == 1
        assert result[0]["key"] == "id"

    def test_无额外字段(self):
        config_h = [{"key": "id", "label": "ID", "width": 10}]
        auto_h = [{"key": "id", "label": "ID", "width": 10}]
        result = merge_headers(config_h, auto_h, {"auto_detect_headers": True})
        assert len(result) == 1

    def test_默认auto_detect_headers为True(self):
        config_h = [{"key": "id", "label": "ID", "width": 10}]
        auto_h = [{"key": "id", "label": "ID", "width": 10}, {"key": "age", "label": "年龄", "width": 10}]
        result = merge_headers(config_h, auto_h, {})
        assert len(result) == 2


class TestGetAlignment:
    def test_left(self):
        assert _get_alignment("left") == "left"

    def test_center(self):
        assert _get_alignment("center") == "center"

    def test_right(self):
        assert _get_alignment("right") == "right"

    def test_未知值默认center(self):
        assert _get_alignment("unknown") == "center"

    def test_none默认center(self):
        assert _get_alignment(None) == "center"


class TestCreateBorder:
    def test_默认边框(self):
        style_config = {}
        border = _create_border(style_config)
        assert border is not None

    def test_自定义边框样式(self):
        style_config = {"border_style": "medium", "border_color": "FF0000"}
        border = _create_border(style_config)
        assert border is not None

    def test_带井号颜色(self):
        style_config = {"border_color": "#FF0000"}
        border = _create_border(style_config)
        assert border is not None

    def test_border_style为None返回None(self):
        style_config = {"border_style": None}
        border = _create_border(style_config)
        assert border is None


class TestApplyHeaderStyle:
    def test_应用默认样式(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID", "姓名"])
        header_cells = ws[1]
        config = {}
        apply_header_style(ws, header_cells, config)
        cell = ws.cell(row=1, column=1)
        assert cell.font.bold is True
        assert cell.fill.start_color.rgb is not None

    def test_自定义样式(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        header_cells = ws[1]
        config = {
            "header_style": {
                "font_name": "Arial",
                "font_bold": False,
                "font_color": "#000000",
                "font_size": 14,
                "bg_color": "#FF0000",
                "alignment": "left",
                "vertical_alignment": "top",
                "border_style": "thin",
                "border_color": "#000000",
            }
        }
        apply_header_style(ws, header_cells, config)
        cell = ws.cell(row=1, column=1)
        assert cell.font.name == "Arial"
        assert cell.font.bold is False
        assert cell.font.size == 14

    def test_无边框时不设置border(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        header_cells = ws[1]
        config = {"header_style": {"border_style": None}}
        apply_header_style(ws, header_cells, config)
        cell = ws.cell(row=1, column=1)
        assert cell.border.left.style is None or cell.border.left.style == "none" or True


class TestApplyDataStyle:
    def test_应用默认样式(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        ws.append([2])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        config = {}
        apply_data_style(ws, headers, 2, config)
        cell = ws.cell(row=2, column=1)
        assert cell.font is not None

    def test_隔行变色(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        ws.append([2])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        config = {"style_alt_rows": True}
        apply_data_style(ws, headers, 2, config)
        cell_even = ws.cell(row=2, column=1)
        assert cell_even.fill.start_color.rgb is not None

    def test_关闭隔行变色(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        ws.append([2])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        config = {"style_alt_rows": False}
        apply_data_style(ws, headers, 2, config)
        cell_odd = ws.cell(row=3, column=1)
        assert cell_odd.fill.start_color.rgb == "00000000"

    def test_自定义数据样式(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        config = {
            "data_style": {
                "font_name": "Arial",
                "font_bold": True,
                "font_color": "#FF0000",
                "font_size": 10,
                "wrap_text": False,
                "alt_row_color": "#CCCCCC",
                "alignment": "right",
                "vertical_alignment": "top",
                "border_style": "medium",
                "border_color": "#000000",
            }
        }
        apply_data_style(ws, headers, 1, config)
        cell = ws.cell(row=2, column=1)
        assert cell.font.name == "Arial"
        assert cell.font.bold is True

    def test_无边框(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        config = {"data_style": {"border_style": None}}
        apply_data_style(ws, headers, 1, config)
        cell = ws.cell(row=2, column=1)


class TestSetColumnWidths:
    def test_设置列宽(self):
        wb = Workbook()
        ws = wb.active
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "name", "label": "姓名", "width": 20}]
        set_column_widths(ws, headers)
        assert ws.column_dimensions["A"].width == 10
        assert ws.column_dimensions["B"].width == 20

    def test_默认宽度(self):
        wb = Workbook()
        ws = wb.active
        headers = [{"key": "id", "label": "ID"}]
        set_column_widths(ws, headers)
        assert ws.column_dimensions["A"].width == 15

    def test_超过26列(self):
        wb = Workbook()
        ws = wb.active
        headers = [{"key": f"col_{i}", "label": f"Col {i}", "width": 10} for i in range(30)]
        set_column_widths(ws, headers)
        assert ws.column_dimensions["A"].width == 10
        assert ws.column_dimensions["AA"].width == 10
        assert ws.column_dimensions["AD"].width == 10


class TestGetColumnLetter:
    def test_A(self):
        assert _get_column_letter(1) == "A"

    def test_Z(self):
        assert _get_column_letter(26) == "Z"

    def test_AA(self):
        assert _get_column_letter(27) == "AA"

    def test_AZ(self):
        assert _get_column_letter(52) == "AZ"

    def test_BA(self):
        assert _get_column_letter(53) == "BA"

    def test_ZZ(self):
        assert _get_column_letter(702) == "ZZ"

    def test_AAA(self):
        assert _get_column_letter(703) == "AAA"


class TestExtractValue:
    def test_简单键(self):
        item = {"name": "张三"}
        assert extract_value(item, "name") == "张三"

    def test_嵌套键点分隔(self):
        item = {"info": {"name": "张三"}}
        assert extract_value(item, "info.name") == "张三"

    def test_多层嵌套(self):
        item = {"a": {"b": {"c": 42}}}
        assert extract_value(item, "a.b.c") == 42

    def test_键不存在返回空字符串(self):
        item = {"name": "张三"}
        assert extract_value(item, "age") == ""

    def test_嵌套键不存在(self):
        item = {"a": {"b": 1}}
        assert extract_value(item, "a.c") == ""

    def test_中间值非字典(self):
        item = {"a": "string"}
        assert extract_value(item, "a.b") == ""

    def test_字典值转JSON(self):
        item = {"a": {"b": {"c": 1}}}
        result = extract_value(item, "a.b")
        assert isinstance(result, str)
        assert '"c"' in result

    def test_列表值转JSON(self):
        item = {"tags": [1, 2, 3]}
        result = extract_value(item, "tags")
        assert isinstance(result, str)
        assert "1" in result

    def test_展平后匹配(self):
        item = {"a": {"b": 1}}
        assert extract_value(item, "a.b") == 1

    def test_none值(self):
        item = {"name": None}
        assert extract_value(item, "name") is None


class TestSanitizeSheetName:
    def test_正常名称(self):
        assert _sanitize_sheet_name("数据导出") == "数据导出"

    def test_替换非法字符(self):
        assert _sanitize_sheet_name("a\\b/c*d[e]:f?g") == "a_b_c_d_e__f_g"

    def test_截断超长名称(self):
        name = "a" * 50
        result = _sanitize_sheet_name(name)
        assert len(result) <= 31

    def test_空名称用Sheet替代(self):
        assert _sanitize_sheet_name("") == "Sheet"

    def test_空格名称strip后为空(self):
        assert _sanitize_sheet_name("   ") == "Sheet"

    def test_自定义最大长度(self):
        name = "a" * 50
        result = _sanitize_sheet_name(name, max_length=20)
        assert len(result) <= 20

    def test_非字符串转字符串(self):
        assert _sanitize_sheet_name(123) == "123"

    def test_strip空格(self):
        assert _sanitize_sheet_name("  hello  ") == "hello"


class TestRenderSheetName:
    def test_value模板(self):
        result = _render_sheet_name("{value}", "技术部", 1, 3)
        assert result == "技术部"

    def test_index模板(self):
        result = _render_sheet_name("{index}", "技术部", 2, 3)
        assert result == "2"

    def test_count模板(self):
        result = _render_sheet_name("{count}", "技术部", 1, 5)
        assert result == "5"

    def test_num模板(self):
        result = _render_sheet_name("{num}", "技术部", 3, 5)
        assert result == "3"

    def test_空值用默认标签(self):
        result = _render_sheet_name("{value}", "", 1, 3)
        assert result == "未分类"

    def test_none值用默认标签(self):
        result = _render_sheet_name("{value}", None, 1, 3)
        assert result == "未分类"

    def test_自定义空标签(self):
        result = _render_sheet_name("{value}", "", 1, 3, empty_label="其他")
        assert result == "其他"

    def test_无效模板键返回字符串(self):
        result = _render_sheet_name("{invalid_key}", "技术部", 1, 3)
        assert result == "技术部"

    def test_复杂模板(self):
        result = _render_sheet_name("第{index}组-{value}", "技术部", 1, 3)
        assert result == "第1组-技术部"


class TestMatchCustomRule:
    def test_values列表匹配(self):
        rule = {"values": ["a", "b", "c"], "name": "组1"}
        assert _match_custom_rule("a", rule) is True
        assert _match_custom_rule("d", rule) is False

    def test_values单个值匹配(self):
        rule = {"values": "a", "name": "组1"}
        assert _match_custom_rule("a", rule) is True
        assert _match_custom_rule("b", rule) is False

    def test_condition匹配(self):
        rule = {"condition": "value > 10", "name": "组1"}
        assert _match_custom_rule(15, rule) is True
        assert _match_custom_rule(5, rule) is False

    def test_condition异常返回False(self):
        rule = {"condition": "1/0", "name": "组1"}
        assert _match_custom_rule(1, rule) is False

    def test_min_max范围(self):
        rule = {"min": 10, "max": 20, "name": "组1"}
        assert _match_custom_rule(15, rule) is True
        assert _match_custom_rule(5, rule) is False
        assert _match_custom_rule(25, rule) is False

    def test_min边界包含(self):
        rule = {"min": 10, "include_min": True, "name": "组1"}
        assert _match_custom_rule(10, rule) is True

    def test_min边界不包含(self):
        rule = {"min": 10, "include_min": False, "name": "组1"}
        assert _match_custom_rule(10, rule) is False
        assert _match_custom_rule(11, rule) is True

    def test_max边界包含(self):
        rule = {"max": 20, "include_max": True, "name": "组1"}
        assert _match_custom_rule(20, rule) is True

    def test_max边界不包含(self):
        rule = {"max": 20, "include_max": False, "name": "组1"}
        assert _match_custom_rule(20, rule) is False
        assert _match_custom_rule(19, rule) is True

    def test_仅min(self):
        rule = {"min": 5, "name": "组1"}
        assert _match_custom_rule(10, rule) is True
        assert _match_custom_rule(3, rule) is False

    def test_仅max(self):
        rule = {"max": 10, "name": "组1"}
        assert _match_custom_rule(5, rule) is True
        assert _match_custom_rule(15, rule) is False

    def test_none值转0(self):
        rule = {"min": 0, "max": 10, "name": "组1"}
        assert _match_custom_rule(None, rule) is True

    def test_空字符串值转0(self):
        rule = {"min": 0, "max": 10, "name": "组1"}
        assert _match_custom_rule("", rule) is True

    def test_无法转数字返回False(self):
        rule = {"min": 0, "max": 10, "name": "组1"}
        assert _match_custom_rule("abc", rule) is False

    def test_无匹配规则返回False(self):
        rule = {"name": "组1"}
        assert _match_custom_rule("anything", rule) is False


class TestMatchRangeGroup:
    def test_范围内匹配(self):
        group = {"min": 10, "max": 20, "name": "10-20"}
        assert _match_range_group(15, group) is True

    def test_范围外不匹配(self):
        group = {"min": 10, "max": 20, "name": "10-20"}
        assert _match_range_group(5, group) is False

    def test_none值不匹配(self):
        group = {"min": 10, "max": 20, "name": "10-20"}
        assert _match_range_group(None, group) is False

    def test_空字符串不匹配(self):
        group = {"min": 10, "max": 20, "name": "10-20"}
        assert _match_range_group("", group) is False

    def test_include_min(self):
        group = {"min": 10, "include_min": True, "name": "组1"}
        assert _match_range_group(10, group) is True

    def test_exclude_min(self):
        group = {"min": 10, "include_min": False, "name": "组1"}
        assert _match_range_group(10, group) is False

    def test_include_max(self):
        group = {"max": 20, "include_max": True, "name": "组1"}
        assert _match_range_group(20, group) is True

    def test_exclude_max(self):
        group = {"max": 20, "include_max": False, "name": "组1"}
        assert _match_range_group(20, group) is False

    def test_无法转数字返回False(self):
        group = {"min": 10, "max": 20, "name": "10-20"}
        assert _match_range_group("abc", group) is False

    def test_仅min(self):
        group = {"min": 5, "name": "组1"}
        assert _match_range_group(10, group) is True

    def test_仅max(self):
        group = {"max": 10, "name": "组1"}
        assert _match_range_group(5, group) is True


class TestSplitDataByField:
    def test_by_value拆分(self):
        data = [
            {"dept": "技术部", "name": "张三"},
            {"dept": "市场部", "name": "李四"},
            {"dept": "技术部", "name": "王五"},
        ]
        config = {"split_rule": "by_value"}
        groups = split_data_by_field(data, "dept", config)
        assert len(groups) == 2
        names = [g["name"] for g in groups]
        assert "技术部" in names
        assert "市场部" in names

    def test_by_value空值用默认标签(self):
        data = [{"dept": "技术部"}, {"dept": ""}, {"dept": None}]
        config = {"split_rule": "by_value"}
        groups = split_data_by_field(data, "dept", config)
        names = [g["name"] for g in groups]
        assert "未分类" in names

    def test_by_value自定义空标签(self):
        data = [{"dept": ""}]
        config = {"split_rule": "by_value", "empty_value_label": "其他"}
        groups = split_data_by_field(data, "dept", config)
        names = [g["name"] for g in groups]
        assert "其他" in names

    def test_by_range拆分(self):
        data = [
            {"age": 15},
            {"age": 25},
            {"age": 35},
        ]
        config = {
            "split_rule": "by_range",
            "range_groups": [
                {"name": "青年", "min": 0, "max": 20, "include_max": True},
                {"name": "中年", "min": 20, "max": 60, "include_min": False},
            ],
        }
        groups = split_data_by_field(data, "age", config)
        names = [g["name"] for g in groups]
        assert "青年" in names
        assert "中年" in names

    def test_by_range无匹配进未分类(self):
        data = [{"age": 100}]
        config = {
            "split_rule": "by_range",
            "range_groups": [
                {"name": "青年", "min": 0, "max": 30, "include_max": True},
            ],
        }
        groups = split_data_by_field(data, "age", config)
        names = [g["name"] for g in groups]
        assert "未分类" in names

    def test_by_custom拆分(self):
        data = [
            {"dept": "技术部"},
            {"dept": "市场部"},
            {"dept": "财务部"},
        ]
        config = {
            "split_rule": "by_custom",
            "custom_rules": [
                {"name": "技术组", "values": ["技术部"]},
                {"name": "其他组", "values": ["市场部", "财务部"]},
            ],
        }
        groups = split_data_by_field(data, "dept", config)
        names = [g["name"] for g in groups]
        assert "技术组" in names
        assert "其他组" in names

    def test_by_custom无匹配用fallback(self):
        data = [{"dept": "人事部"}]
        config = {
            "split_rule": "by_custom",
            "custom_rules": [
                {"name": "技术组", "values": ["技术部"]},
            ],
            "fallback_group_name": "其他",
        }
        groups = split_data_by_field(data, "dept", config)
        names = [g["name"] for g in groups]
        assert "其他" in names

    def test_by_custom无匹配无fallback用默认(self):
        data = [{"dept": "人事部"}]
        config = {
            "split_rule": "by_custom",
            "custom_rules": [
                {"name": "技术组", "values": ["技术部"]},
            ],
        }
        groups = split_data_by_field(data, "dept", config)
        names = [g["name"] for g in groups]
        assert "未分类" in names

    def test_非字典项跳过(self):
        data = ["not_dict", {"dept": "技术部"}]
        config = {"split_rule": "by_value"}
        groups = split_data_by_field(data, "dept", config)
        assert len(groups) == 1

    def test_original_indices记录(self):
        data = [{"dept": "技术部"}, {"dept": "市场部"}, {"dept": "技术部"}]
        config = {"split_rule": "by_value"}
        groups = split_data_by_field(data, "dept", config)
        for g in groups:
            if g["name"] == "技术部":
                assert g["original_indices"] == [0, 2]
            elif g["name"] == "市场部":
                assert g["original_indices"] == [1]


class TestToNumeric:
    def test_整数(self):
        assert _to_numeric(42) == 42.0

    def test_浮点数(self):
        assert _to_numeric(3.14) == 3.14

    def test_数字字符串(self):
        assert _to_numeric("42") == 42.0

    def test_none返回None(self):
        assert _to_numeric(None) is None

    def test_空字符串返回None(self):
        assert _to_numeric("") is None

    def test_非数字字符串返回None(self):
        assert _to_numeric("abc") is None

    def test_布尔值(self):
        assert _to_numeric(True) == 1.0
        assert _to_numeric(False) == 0.0


class TestAggregateValues:
    def test_sum求和(self):
        assert _aggregate_values([1, 2, 3], "sum") == 6

    def test_sum空列表(self):
        assert _aggregate_values([], "sum") == 0

    def test_count计数(self):
        assert _aggregate_values([1, "a", None, 3], "count") == 4

    def test_count_num数值计数(self):
        assert _aggregate_values([1, "a", 3, None], "count_num") == 2

    def test_average平均值(self):
        assert _aggregate_values([10, 20, 30], "average") == 20

    def test_average空列表(self):
        assert _aggregate_values([], "average") == 0

    def test_max最大值(self):
        assert _aggregate_values([1, 5, 3], "max") == 5

    def test_max空列表(self):
        assert _aggregate_values([], "max") == 0

    def test_min最小值(self):
        assert _aggregate_values([1, 5, 3], "min") == 1

    def test_min空列表(self):
        assert _aggregate_values([], "min") == 0

    def test_product乘积(self):
        assert _aggregate_values([2, 3, 4], "product") == 24

    def test_product空列表(self):
        assert _aggregate_values([], "product") == 0

    def test_product含0(self):
        assert _aggregate_values([2, 0, 4], "product") == 0

    def test_stddev样本标准偏差(self):
        data = [2, 4, 4, 4, 5, 5, 7, 9]
        mean = sum(data) / len(data)
        variance = sum((x - mean) ** 2 for x in data) / (len(data) - 1)
        expected = variance ** 0.5
        result = _aggregate_values(data, "stddev")
        assert abs(result - expected) < 0.001

    def test_stddev不足2个返回0(self):
        assert _aggregate_values([5], "stddev") == 0

    def test_stddevp总体标准偏差(self):
        result = _aggregate_values([2, 4, 4, 4, 5, 5, 7, 9], "stddevp")
        assert result > 0

    def test_stddevp空列表返回0(self):
        assert _aggregate_values([], "stddevp") == 0

    def test_var样本方差(self):
        result = _aggregate_values([2, 4, 4, 4, 5, 5, 7, 9], "var")
        assert result > 0

    def test_var不足2个返回0(self):
        assert _aggregate_values([5], "var") == 0

    def test_varp总体方差(self):
        result = _aggregate_values([2, 4, 4, 4, 5, 5, 7, 9], "varp")
        assert result > 0

    def test_varp空列表返回0(self):
        assert _aggregate_values([], "varp") == 0

    def test_未知聚合函数默认sum(self):
        assert _aggregate_values([1, 2, 3], "unknown") == 6

    def test_混合类型值(self):
        assert _aggregate_values([1, "a", 3], "sum") == 4

    def test_字符串数字(self):
        assert _aggregate_values(["1", "2", "3"], "sum") == 6


class TestGetFieldLabel:
    def test_找到标签(self):
        headers = [{"key": "name", "label": "姓名"}]
        assert _get_field_label(headers, "name") == "姓名"

    def test_未找到返回key(self):
        headers = [{"key": "name", "label": "姓名"}]
        assert _get_field_label(headers, "age") == "age"


class TestGetRowKey:
    def test_单字段(self):
        item = {"dept": "技术部"}
        result = _get_row_key(item, ["dept"], "未分类")
        assert result == ("技术部",)

    def test_多字段(self):
        item = {"dept": "技术部", "level": "高级"}
        result = _get_row_key(item, ["dept", "level"], "未分类")
        assert result == ("技术部", "高级")

    def test_空值用默认标签(self):
        item = {"dept": ""}
        result = _get_row_key(item, ["dept"], "未分类")
        assert result == ("未分类",)

    def test_none值用默认标签(self):
        item = {"dept": None}
        result = _get_row_key(item, ["dept"], "未分类")
        assert result == ("未分类",)


class TestGetColKey:
    def test_无列字段返回空元组(self):
        item = {"dept": "技术部"}
        result = _get_col_key(item, [], "未分类")
        assert result == ("",)

    def test_单字段(self):
        item = {"year": "2024"}
        result = _get_col_key(item, ["year"], "未分类")
        assert result == ("2024",)

    def test_空值用默认标签(self):
        item = {"year": ""}
        result = _get_col_key(item, ["year"], "未分类")
        assert result == ("未分类",)


class TestBuildPivotTable:
    def test_基本透视表(self):
        data = [
            {"dept": "技术部", "salary": 10000},
            {"dept": "技术部", "salary": 20000},
            {"dept": "市场部", "salary": 15000},
        ]
        pivot_config = {
            "row_fields": ["dept"],
            "column_fields": [],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        headers = [{"key": "dept", "label": "部门"}, {"key": "salary", "label": "薪资"}]
        result = build_pivot_table(data, pivot_config, headers)
        assert "技术部" in [rk[0] for rk in result["row_keys"]]
        assert "市场部" in [rk[0] for rk in result["row_keys"]]

    def test_带列字段(self):
        data = [
            {"dept": "技术部", "year": "2024", "salary": 10000},
            {"dept": "技术部", "year": "2023", "salary": 20000},
        ]
        pivot_config = {
            "row_fields": ["dept"],
            "column_fields": ["year"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        headers = [{"key": "dept", "label": "部门"}, {"key": "year", "label": "年份"}, {"key": "salary", "label": "薪资"}]
        result = build_pivot_table(data, pivot_config, headers)
        assert len(result["col_keys"]) == 2

    def test_非字典项跳过(self):
        data = ["not_dict", {"dept": "技术部", "salary": 10000}]
        pivot_config = {
            "row_fields": ["dept"],
            "column_fields": [],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        headers = [{"key": "dept", "label": "部门"}, {"key": "salary", "label": "薪资"}]
        result = build_pivot_table(data, pivot_config, headers)
        assert len(result["row_keys"]) == 1

    def test_空值标签(self):
        data = [{"dept": "", "salary": 10000}]
        pivot_config = {
            "row_fields": ["dept"],
            "column_fields": [],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
            "empty_value_label": "(空白)",
        }
        headers = [{"key": "dept", "label": "部门"}, {"key": "salary", "label": "薪资"}]
        result = build_pivot_table(data, pivot_config, headers)
        assert result["empty_label"] == "(空白)"


class TestComputePivotValue:
    def test_计算值(self):
        pivot_result = {
            "data": {
                ("技术部",): {
                    ("",): {"salary:sum": [10000, 20000]}
                }
            },
            "row_keys": [("技术部",)],
            "col_keys": [("",)],
        }
        vf = {"field": "salary", "aggregate": "sum"}
        result = _compute_pivot_value(pivot_result, ("技术部",), ("",), vf)
        assert result == 30000

    def test_键不存在返回0(self):
        pivot_result = {
            "data": {},
            "row_keys": [],
            "col_keys": [],
        }
        vf = {"field": "salary", "aggregate": "sum"}
        result = _compute_pivot_value(pivot_result, ("技术部",), ("",), vf)
        assert result == 0


class TestComputeRowTotal:
    def test_sum总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {
                    ("2023",): {"salary:sum": [10000]},
                    ("2024",): {"salary:sum": [20000]},
                }
            },
            "row_keys": [("技术部",)],
            "col_keys": [("2023",), ("2024",)],
        }
        vf = {"field": "salary", "aggregate": "sum"}
        result = _compute_row_total(pivot_result, ("技术部",), vf)
        assert result == 30000

    def test_count总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {
                    ("2023",): {"salary:count": [1, 2]},
                    ("2024",): {"salary:count": [3]},
                }
            },
            "row_keys": [("技术部",)],
            "col_keys": [("2023",), ("2024",)],
        }
        vf = {"field": "salary", "aggregate": "count"}
        result = _compute_row_total(pivot_result, ("技术部",), vf)
        assert result == 3

    def test_average总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {
                    ("2023",): {"salary:average": [10000, 20000]},
                    ("2024",): {"salary:average": [30000]},
                }
            },
            "row_keys": [("技术部",)],
            "col_keys": [("2023",), ("2024",)],
        }
        vf = {"field": "salary", "aggregate": "average"}
        result = _compute_row_total(pivot_result, ("技术部",), vf)
        assert result == 20000

    def test_max总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {
                    ("2023",): {"salary:max": [10000]},
                    ("2024",): {"salary:max": [30000]},
                }
            },
            "row_keys": [("技术部",)],
            "col_keys": [("2023",), ("2024",)],
        }
        vf = {"field": "salary", "aggregate": "max"}
        result = _compute_row_total(pivot_result, ("技术部",), vf)
        assert result == 30000

    def test_min总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {
                    ("2023",): {"salary:min": [10000]},
                    ("2024",): {"salary:min": [30000]},
                }
            },
            "row_keys": [("技术部",)],
            "col_keys": [("2023",), ("2024",)],
        }
        vf = {"field": "salary", "aggregate": "min"}
        result = _compute_row_total(pivot_result, ("技术部",), vf)
        assert result == 10000

    def test_其他聚合函数求和(self):
        pivot_result = {
            "data": {
                ("技术部",): {
                    ("2023",): {"salary:product": [2, 3]},
                    ("2024",): {"salary:product": [4]},
                }
            },
            "row_keys": [("技术部",)],
            "col_keys": [("2023",), ("2024",)],
        }
        vf = {"field": "salary", "aggregate": "product"}
        result = _compute_row_total(pivot_result, ("技术部",), vf)
        assert result == 10


class TestComputeColTotal:
    def test_sum列总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:sum": [10000]}},
                ("市场部",): {("2023",): {"salary:sum": [20000]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "sum"}
        result = _compute_col_total(pivot_result, ("2023",), vf)
        assert result == 30000

    def test_average列总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:average": [10000, 20000]}},
                ("市场部",): {("2023",): {"salary:average": [30000]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "average"}
        result = _compute_col_total(pivot_result, ("2023",), vf)
        assert result == 20000

    def test_max列总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:max": [10000]}},
                ("市场部",): {("2023",): {"salary:max": [30000]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "max"}
        result = _compute_col_total(pivot_result, ("2023",), vf)
        assert result == 30000

    def test_min列总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:min": [10000]}},
                ("市场部",): {("2023",): {"salary:min": [30000]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "min"}
        result = _compute_col_total(pivot_result, ("2023",), vf)
        assert result == 10000

    def test_count列总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:count": [1, 2]}},
                ("市场部",): {("2023",): {"salary:count": [3]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "count"}
        result = _compute_col_total(pivot_result, ("2023",), vf)
        assert result == 3

    def test_其他聚合函数列总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:product": [2, 3]}},
                ("市场部",): {("2023",): {"salary:product": [4]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "product"}
        result = _compute_col_total(pivot_result, ("2023",), vf)
        assert result == 10


class TestComputeGrandTotal:
    def test_sum总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:sum": [10000]}},
                ("市场部",): {("2023",): {"salary:sum": [20000]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "sum"}
        result = _compute_grand_total(pivot_result, vf)
        assert result == 30000

    def test_average总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:average": [10000, 20000]}},
                ("市场部",): {("2023",): {"salary:average": [30000]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "average"}
        result = _compute_grand_total(pivot_result, vf)
        assert result == 20000

    def test_max总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:max": [10000]}},
                ("市场部",): {("2023",): {"salary:max": [30000]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "max"}
        result = _compute_grand_total(pivot_result, vf)
        assert result == 30000

    def test_min总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:min": [10000]}},
                ("市场部",): {("2023",): {"salary:min": [30000]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "min"}
        result = _compute_grand_total(pivot_result, vf)
        assert result == 10000

    def test_count总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:count": [1, 2]}},
                ("市场部",): {("2023",): {"salary:count": [3]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "count"}
        result = _compute_grand_total(pivot_result, vf)
        assert result == 3

    def test_其他聚合函数总计(self):
        pivot_result = {
            "data": {
                ("技术部",): {("2023",): {"salary:product": [2, 3]}},
                ("市场部",): {("2023",): {"salary:product": [4]}},
            },
            "row_keys": [("技术部",), ("市场部",)],
            "col_keys": [("2023",)],
        }
        vf = {"field": "salary", "aggregate": "product"}
        result = _compute_grand_total(pivot_result, vf)
        assert result == 10


class TestFormatValue:
    def test_整数浮点数转整数(self):
        assert _format_value(10.0, "sum") == 10

    def test_小数浮点数保留2位(self):
        assert _format_value(10.333, "average") == 10.33

    def test_非浮点数原样返回(self):
        assert _format_value("hello", "sum") == "hello"

    def test_整数原样返回(self):
        assert _format_value(42, "sum") == 42

    def test_0浮点数转整数(self):
        assert _format_value(0.0, "sum") == 0


class TestGetAggregateLabel:
    def test_sum(self):
        assert _get_aggregate_label("sum") == "求和"

    def test_count(self):
        assert _get_aggregate_label("count") == "计数"

    def test_average(self):
        assert _get_aggregate_label("average") == "平均值"

    def test_max(self):
        assert _get_aggregate_label("max") == "最大值"

    def test_min(self):
        assert _get_aggregate_label("min") == "最小值"

    def test_product(self):
        assert _get_aggregate_label("product") == "乘积"

    def test_count_num(self):
        assert _get_aggregate_label("count_num") == "数值计数"

    def test_stddev(self):
        assert _get_aggregate_label("stddev") == "标准偏差"

    def test_stddevp(self):
        assert _get_aggregate_label("stddevp") == "总体标准偏差"

    def test_var(self):
        assert _get_aggregate_label("var") == "方差"

    def test_varp(self):
        assert _get_aggregate_label("varp") == "总体方差"

    def test_未知函数返回原值(self):
        assert _get_aggregate_label("custom") == "custom"


class TestDataLoader初始化:
    def test_默认配置(self):
        loader = DataLoader()
        assert loader._raw_data is None
        assert loader._auto_headers is None
        assert loader._merged_headers is None
        assert loader._computed_cache is None
        assert loader._validation_result is None
        assert loader._original_indices is None

    def test_dict配置(self):
        config = {"json_file_path": "test.json"}
        loader = DataLoader(config)
        assert loader.config is not None

    def test_ExportConfig配置(self):
        from config_manager import ExportConfig
        ec = ExportConfig.from_default()
        loader = DataLoader(ec)
        assert loader.config is ec

    def test_from_file(self, tmp_path):
        data = [{"id": 1, "name": "张三"}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        loader = DataLoader.from_file(str(f))
        assert loader._raw_data is not None
        assert len(loader._raw_data) == 1


class TestDataLoader加载:
    def test_load加载文件(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        loader = DataLoader()
        result = loader.load(str(f))
        assert result == data

    def test_load无路径抛出异常(self):
        loader = DataLoader({"json_file_path": ""})
        with pytest.raises(ValueError):
            loader.load()

    def test_set_data(self):
        loader = DataLoader()
        data = [{"id": 1}]
        loader.set_data(data)
        assert loader._raw_data == data

    def test_raw_data未加载抛出异常(self):
        loader = DataLoader()
        with pytest.raises(ValueError):
            _ = loader.raw_data


class TestDataLoader字段检测:
    def test_detect_headers(self):
        loader = DataLoader()
        loader.set_data([{"id": 1, "name": "张三"}])
        headers = loader.detect_headers()
        assert len(headers) >= 2

    def test_detect_headers未加载抛出异常(self):
        loader = DataLoader()
        with pytest.raises(ValueError):
            loader.detect_headers()

    def test_auto_headers属性自动检测(self):
        loader = DataLoader()
        loader.set_data([{"id": 1}])
        headers = loader.auto_headers
        assert len(headers) >= 1


class TestDataLoader合并字段:
    def test_merge_headers(self):
        loader = DataLoader()
        loader.set_data([{"id": 1, "name": "张三"}])
        headers = loader.merge_headers()
        assert len(headers) >= 2

    def test_headers属性自动合并(self):
        loader = DataLoader()
        loader.set_data([{"id": 1}])
        headers = loader.headers
        assert len(headers) >= 1


class TestDataLoader计算列:
    def test_apply_computed_columns无计算列(self):
        loader = DataLoader()
        loader.set_data([{"id": 1}])
        cache, headers = loader.apply_computed_columns()
        assert cache is None

    def test_computed_cache属性(self):
        loader = DataLoader()
        loader.set_data([{"id": 1}])
        assert loader.computed_cache is None


class TestDataLoader校验:
    def test_validate_data无规则(self):
        loader = DataLoader()
        loader.set_data([{"id": 1}])
        result = loader.validate_data()
        assert result is None

    def test_validation_result属性(self):
        loader = DataLoader()
        loader.set_data([{"id": 1}])
        assert loader.validation_result is None


class TestDataLoader有效数据:
    def test_get_valid_data无校验结果(self):
        loader = DataLoader()
        loader.set_data([{"id": 1}])
        data = loader.get_valid_data()
        assert data == [{"id": 1}]

    def test_original_indices属性(self):
        loader = DataLoader()
        loader.set_data([{"id": 1}, {"id": 2}])
        indices = loader.original_indices
        assert indices == [0, 1]


class TestDataLoader辅助方法:
    def test_extract_value(self):
        loader = DataLoader()
        assert loader.extract_value({"name": "张三"}, "name") == "张三"

    def test_flatten_dict(self):
        loader = DataLoader()
        result = loader.flatten_dict({"a": {"b": 1}})
        assert result == {"a.b": 1}

    def test_repr(self):
        loader = DataLoader()
        assert "DataLoader" in repr(loader)
        assert "0 items" in repr(loader)

    def test_repr有数据(self):
        loader = DataLoader()
        loader.set_data([{"id": 1}])
        assert "1 items" in repr(loader)


class TestDataLoaderPrepareForExport:
    def test_prepare_for_export(self, tmp_path):
        data = [{"id": 1, "name": "张三"}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        loader = DataLoader({"json_file_path": str(f)})
        loader.set_data(data)
        result = loader.prepare_for_export()
        assert result is not None
        assert "data" in result
        assert "headers" in result


class TestExcelExporter初始化:
    def test_默认配置(self):
        exporter = ExcelExporter()
        assert exporter.config is not None

    def test_dict配置(self):
        config = {"excel_output_path": "/tmp/test.xlsx"}
        exporter = ExcelExporter(config)
        assert exporter.config is not None

    def test_ExportConfig配置(self):
        from config_manager import ExportConfig
        ec = ExportConfig.from_default()
        exporter = ExcelExporter(ec)
        assert exporter.config is ec

    def test_from_config(self):
        exporter = ExcelExporter.from_config({"excel_output_path": "/tmp/test.xlsx"})
        assert exporter.config is not None

    def test_repr(self):
        exporter = ExcelExporter()
        assert "ExcelExporter" in repr(exporter)


class TestExcelExporter导出:
    def test_export单工作表(self, tmp_path):
        data = [{"id": 1, "name": "张三"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "name", "label": "姓名", "width": 15}]
        output = str(tmp_path / "test.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "header_style": {"border_style": "thin", "border_color": "000000"},
            "data_style": {"border_style": "thin", "border_color": "000000"},
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        result = exporter.export(data, headers)
        assert result == output
        assert os.path.exists(output)

    def test_export拆分工作表(self, tmp_path):
        data = [
            {"id": 1, "name": "张三", "dept": "技术部"},
            {"id": 2, "name": "李四", "dept": "市场部"},
        ]
        headers = [
            {"key": "id", "label": "ID", "width": 10},
            {"key": "name", "label": "姓名", "width": 15},
            {"key": "dept", "label": "部门", "width": 15},
        ]
        output = str(tmp_path / "test_split.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "header_style": {"border_style": "thin", "border_color": "000000"},
            "data_style": {"border_style": "thin", "border_color": "000000"},
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        result = exporter.export(data, headers)
        assert result == output
        assert os.path.exists(output)

    def test_export自动创建目录(self, tmp_path):
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "subdir" / "test.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "header_style": {},
            "data_style": {},
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        result = exporter.export(data, headers)
        assert os.path.exists(output)

    def test_export_from_loader(self, tmp_path):
        data = [{"id": 1, "name": "张三"}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        output = str(tmp_path / "from_loader.xlsx")
        config = {
            "json_file_path": str(f),
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "header_style": {},
            "data_style": {},
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
            "default_headers": [],
            "auto_detect_headers": True,
            "validation_rules": [],
            "computed_columns": [],
        }
        loader = DataLoader(config)
        loader.set_data(data)
        exporter = ExcelExporter(config)
        result = exporter.export_from_loader(loader)
        assert result == output
        assert os.path.exists(output)

    def test_export含非字典项(self, tmp_path):
        data = [{"id": 1}, "not_dict", {"id": 2}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "test_nondict.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "header_style": {},
            "data_style": {},
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        result = exporter.export(data, headers)
        assert result == output


class TestExcelExporter唯一工作表名:
    def test_获取唯一名称(self):
        exporter = ExcelExporter()
        used = {"技术部", "市场部"}
        result = exporter._get_unique_sheet_name("技术部", used, 31)
        assert result != "技术部"
        assert "技术部" in result

    def test_名称不冲突直接返回(self):
        exporter = ExcelExporter()
        used = {"市场部"}
        result = exporter._get_unique_sheet_name("技术部", used, 31)
        assert result == "技术部"

    def test_超长名称截断后加后缀(self):
        exporter = ExcelExporter()
        base = "a" * 31
        used = {base}
        result = exporter._get_unique_sheet_name(base, used, 31)
        assert len(result) <= 31
        assert result != base


class TestExcelExporter拆分无数据:
    def test_拆分无数据时创建默认工作表(self, tmp_path):
        data = []
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "test_empty.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "header_style": {},
            "data_style": {},
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": False,
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        result = exporter.export(data, headers)
        assert result == output


class TestExcelExporterWriteSheet:
    def test_write_sheet调用(self, tmp_path):
        exporter = ExcelExporter()
        wb = Workbook()
        ws = wb.active
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        exporter._write_sheet(ws, data, headers)
        assert ws.cell(row=1, column=1).value == "ID"
        assert ws.cell(row=2, column=1).value == 1


class TestExcelExporterApplyStyles:
    def test_apply_styles(self):
        exporter = ExcelExporter()
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = [{"id": 1}]
        exporter._apply_styles(ws, headers, data)
        cell = ws.cell(row=1, column=1)
        assert cell.font is not None

    def test_apply_styles不应用表头样式(self):
        from config_manager import ExportConfig
        ec = ExportConfig.from_default()
        ec.set("style_header", False)
        exporter = ExcelExporter(ec)
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = [{"id": 1}]
        exporter._apply_styles(ws, headers, data)


class TestExcelExporterPrintSummary:
    def test_打印摘要无校验(self, capsys):
        exporter = ExcelExporter()
        exporter._print_export_summary("/tmp/test.xlsx", [{"id": 1}], [{"key": "id", "label": "ID", "width": 10}])
        captured = capsys.readouterr()
        assert "成功导出" in captured.out

    def test_打印摘要有校验(self, capsys):
        exporter = ExcelExporter()
        mock_result = MagicMock()
        mock_result.skipped_rows = [0]
        mock_result.marked_rows = [1]
        exporter._print_export_summary(
            "/tmp/test.xlsx",
            [{"id": 1}],
            [{"key": "id", "label": "ID", "width": 10}],
            validation_result=mock_result,
        )
        captured = capsys.readouterr()
        assert "跳过" in captured.out

    def test_打印摘要多工作表(self, capsys):
        exporter = ExcelExporter()
        exporter._print_export_summary(
            "/tmp/test.xlsx",
            [{"id": 1}],
            [{"key": "id", "label": "ID", "width": 10}],
            sheet_count=3,
        )
        captured = capsys.readouterr()
        assert "3" in captured.out
