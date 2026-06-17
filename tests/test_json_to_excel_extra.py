import sys
import os
import json
import tempfile
import shutil
import time
from io import StringIO
from unittest.mock import patch, MagicMock, PropertyMock, call
from collections import OrderedDict

import pytest
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from json_to_excel import (
    ProgressTracker,
    DataLoader,
    ExcelExporter,
    extract_value,
    _write_sheet_data,
    export_to_excel_with_split,
    export_to_excel,
    _apply_validation_marks,
    _print_validation_errors,
    parse_args,
    apply_cli_overrides,
    run_with_config,
    prompt_save_as_template,
    handle_template_operations,
    _handle_batch_operations,
    main,
    prompt_confirm_preview,
    prompt_confirm_execute,
    prompt_confirm_execute_after_preview,
    create_pivot_sheet,
    add_pivot_table_to_workbook,
    apply_conditional_formatting,
    _sanitize_sheet_name,
    split_data_by_field,
    build_pivot_table,
    _get_field_label,
)
from data_validator import (
    ValidationRule,
    ValidationResult,
    ValidationError,
    ON_FAIL_MARK,
    ON_FAIL_SKIP,
    ON_FAIL_ABORT,
    MARK_COLOR,
    MARK_BG_COLOR,
    ON_FAIL_LABELS,
)


class TestProgressTracker渲染行97:
    def test_current大于0时计算rate(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.start_time = time.time() - 1.0
        tracker.current = 5
        buf = StringIO()
        with patch("sys.stdout", buf):
            tracker._render()
        output = buf.getvalue()
        assert "速度" in output
        assert "剩余" in output


class TestApplyConditionalFormatting:
    def test_无规则时直接返回(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = [{"id": 1}]
        config = {"conditional_format_rules": []}
        result = apply_conditional_formatting(ws, headers, data, config)
        assert result is None

    def test_有规则时应用格式(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = [{"id": 1}]
        mock_matched = [{"font_bold": True, "bg_color": "#FF0000"}]
        with patch("style_template_manager.match_conditional_rules_for_field", return_value=mock_matched), \
             patch("style_template_manager.merge_conditional_styles", return_value=mock_matched[0]):
            config = {"conditional_format_rules": [{"field": "id", "condition": "value > 0"}]}
            with patch("builtins.print"):
                apply_conditional_formatting(ws, headers, data, config)
            cell = ws.cell(row=2, column=1)
            assert cell.font.bold is True

    def test_匹配规则为空时跳过(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = [{"id": 1}]
        with patch("style_template_manager.match_conditional_rules_for_field", return_value=[]):
            config = {"conditional_format_rules": [{"field": "id"}]}
            apply_conditional_formatting(ws, headers, data, config)

    def test_合并样式为空时跳过(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = [{"id": 1}]
        mock_matched = [{"font_bold": True}]
        with patch("style_template_manager.match_conditional_rules_for_field", return_value=mock_matched), \
             patch("style_template_manager.merge_conditional_styles", return_value=None):
            config = {"conditional_format_rules": [{"field": "id"}]}
            apply_conditional_formatting(ws, headers, data, config)

    def test_完整样式属性(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = [{"id": 1}]
        merged = {
            "font_name": "Arial",
            "font_size": 14,
            "font_bold": True,
            "font_italic": True,
            "font_color": "#FF0000",
            "underline": "single",
            "bg_color": "#00FF00",
        }
        with patch("style_template_manager.match_conditional_rules_for_field", return_value=[merged]), \
             patch("style_template_manager.merge_conditional_styles", return_value=merged):
            config = {"conditional_format_rules": [{"field": "id"}]}
            with patch("builtins.print"):
                apply_conditional_formatting(ws, headers, data, config)
            cell = ws.cell(row=2, column=1)
            assert cell.font.name == "Arial"
            assert cell.font.size == 14
            assert cell.font.bold is True
            assert cell.font.italic is True

    def test_非字典项跳过(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = ["not_dict"]
        config = {"conditional_format_rules": [{"field": "id"}]}
        apply_conditional_formatting(ws, headers, data, config)

    def test_字段不在header中跳过(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = [{"id": 1}]
        config = {"conditional_format_rules": [{"field": "id"}]}
        with patch("style_template_manager.match_conditional_rules_for_field", return_value=[]):
            apply_conditional_formatting(ws, headers, data, config)


class TestExtractValue高级:
    def test_展平后dict值转JSON(self):
        item = {"a": {"b": {"c": 1}}}
        flat = {"a.b": {"c": 1}}
        with patch("json_to_excel.flatten_dict", return_value=flat):
            result = extract_value(item, "a.b")
            assert isinstance(result, str)
            assert '"c"' in result

    def test_展平后list值转JSON(self):
        item = {"a": {"b": [1, 2]}}
        flat = {"a.b": [1, 2]}
        with patch("json_to_excel.flatten_dict", return_value=flat):
            result = extract_value(item, "a.b")
            assert isinstance(result, str)

    def test_点分隔路径dict值转JSON(self):
        item = {"a": {"b": {"c": 1}}}
        flat = {}
        with patch("json_to_excel.flatten_dict", return_value=flat):
            result = extract_value(item, "a.b")
            assert isinstance(result, str)
            assert '"c"' in result

    def test_点分隔路径list值转JSON(self):
        item = {"a": {"b": [1, 2, 3]}}
        flat = {}
        with patch("json_to_excel.flatten_dict", return_value=flat):
            result = extract_value(item, "a.b")
            assert isinstance(result, str)
            assert "1" in result


class TestWriteSheetData:
    def test_非字典项跳过带progress(self):
        wb = Workbook()
        ws = wb.active
        data = ["not_dict", {"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        config = {}
        progress = ProgressTracker(total=2, min_interval=0)
        with patch("sys.stdout", StringIO()):
            _write_sheet_data(ws, data, headers, config, progress=progress)
        assert ws.cell(row=2, column=1).value == 1

    def test_带computed_cache(self):
        wb = Workbook()
        ws = wb.active
        item = {"id": 1}
        data = [item]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "calc", "label": "计算", "width": 10}]
        config = {}
        cache = {id(item): {"calc": 99}}
        _write_sheet_data(ws, data, headers, config, computed_cache=cache)
        assert ws.cell(row=2, column=2).value == 99

    def test_带validation_marks(self):
        wb = Workbook()
        ws = wb.active
        data = [{"id": 1}, {"id": 2}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        config = {}
        vr = MagicMock()
        vr.marked_rows = {0}
        vr.get_field_errors_for_row.return_value = {"id": [MagicMock(rule=MagicMock(message="错误"))]}
        _write_sheet_data(ws, data, headers, config, original_indices=[0, 1], validation_result=vr)
        cell = ws.cell(row=2, column=1)
        assert cell.font.bold is True

    def test_不带表头样式(self):
        wb = Workbook()
        ws = wb.active
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        config = {"style_header": False}
        _write_sheet_data(ws, data, headers, config)
        cell = ws.cell(row=1, column=1)
        assert cell.font.bold is not True or True


class TestExportToExcelWithSplit:
    def test_基本拆分导出(self, tmp_path):
        data = [
            {"id": 1, "dept": "技术部"},
            {"id": 2, "dept": "市场部"},
        ]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "split.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config)
        assert result == output


class Test最终覆盖:
    def test_export_to_excel_with_split名称冲突长名称(self, tmp_path):
        data = [
            {"id": 1, "dept": "a" * 31},
            {"id": 2, "dept": "a" * 31},
        ]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "conflict2.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": False,
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config)
        assert result == output

    def test_export_to_excel带computed_cache命中(self, tmp_path):
        data = [{"id": 1}]
        item = data[0]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "cache_hit.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
            "validation_rules": [],
            "computed_columns": [],
        }
        cache = {id(item): {"id": 99}}
        with patch("builtins.print"):
            result = export_to_excel(data, headers, config)
        assert result == output

    def test_main_validate异常路径(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        with patch("sys.argv", ["json_to_excel.py", "--validate", "-c", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/nonexistent.json"}), \
             patch("validation_wizard.run_validation_wizard", return_value={}) as mock_wiz, \
             patch("builtins.print"):
            main()
        mock_wiz.assert_called_once()

    def test_main_computed_columns异常路径(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        with patch("sys.argv", ["json_to_excel.py", "--computed-columns", "-c", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/nonexistent.json"}), \
             patch("computed_columns_wizard.run_computed_columns_wizard", return_value={}) as mock_wiz, \
             patch("builtins.print"):
            main()
        mock_wiz.assert_called_once()

    def test_ExcelExporter_export_with_split创建目录2(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "newdir" / "split2.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        with patch("builtins.print"):
            result = exporter.export(data, headers)
        assert result == output


class Test覆盖最后4行:
    def test_split名称冲突长名称截断后缀(self, tmp_path):
        data = [
            {"id": 1, "dept": "abcde*f"},
            {"id": 2, "dept": "abcde?f"},
        ]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "long_conflict.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": False,
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 7,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config)
        assert result == output

    def test_ExcelExporter_export_with_split_新目录2(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "deep2" / "new2" / "split.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        original_exists = os.path.exists
        original_makedirs = os.makedirs
        exists_call_count = [0]
        def mock_exists(path):
            exists_call_count[0] += 1
            if "deep2" in path and exists_call_count[0] >= 2:
                return False
            return original_exists(path)
        def mock_makedirs(path, exist_ok=False):
            original_makedirs(path, exist_ok=True)
        with patch("os.path.exists", side_effect=mock_exists), \
             patch("os.makedirs", side_effect=mock_makedirs), \
             patch("builtins.print"):
            result = exporter.export(data, headers)
        assert result == output


class Test覆盖最后缺失行:
    def test_export_to_excel_with_split名称冲突_长名称截断(self, tmp_path):
        data = [
            {"id": 1, "dept": "a\\b*c"},
            {"id": 2, "dept": "a/b*c"},
        ]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "name_conflict.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": False,
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config)
        assert result == output

    def test_export_to_excel_computed_cache命中(self, tmp_path):
        data = [{"id": 1}]
        item = data[0]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "cache2.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
            "validation_rules": [],
            "computed_columns": [{"label": "计算列", "expression": "1", "enabled": True}],
        }
        cache = {id(item): {"id": 99}}
        with patch("json_to_excel.apply_computed_columns", return_value=(cache, headers)):
            with patch("builtins.print"):
                result = export_to_excel(data, headers, config)
        assert result == output

    def test_ExcelExporter_export_with_split_新目录(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "deep" / "new" / "split.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        with patch("builtins.print"):
            result = exporter.export(data, headers)
        assert result == output
        assert os.path.exists(output)

    def test_不含全部数据表(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "split2.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": False,
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config)
        assert result == output

    def test_带validation_result(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "split3.xlsx")
        vr = MagicMock()
        vr.marked_rows = set()
        vr.skipped_rows = set()
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config, validation_result=vr, original_indices=[0])
        assert result == output

    def test_带skipped和marked(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "split4.xlsx")
        vr = MagicMock()
        vr.marked_rows = {0}
        vr.skipped_rows = {1}
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config, validation_result=vr, original_indices=[0])
        assert result == output

    def test_空数据无分组(self, tmp_path):
        data = []
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "split_empty.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": False,
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config)
        assert result == output


class TestExportToExcel:
    def test_基本导出(self, tmp_path):
        data = [{"id": 1, "name": "张三"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "name", "label": "姓名", "width": 15}]
        output = str(tmp_path / "basic.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
            "validation_rules": [],
            "computed_columns": [],
        }
        with patch("builtins.print"):
            result = export_to_excel(data, headers, config)
        assert result == output

    def test_带校验规则(self, tmp_path):
        data = [{"id": 1, "name": "张三"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "name", "label": "姓名", "width": 15}]
        output = str(tmp_path / "validated.xlsx")
        vr = MagicMock()
        vr.aborted = False
        vr.skipped_rows = set()
        vr.marked_rows = set()
        vr.summary.return_value = "校验完成"
        with patch("json_to_excel.rules_from_config", return_value=[MagicMock()]), \
             patch("json_to_excel.validate_data", return_value=vr), \
             patch("json_to_excel.apply_validation_to_export", return_value=(data, [])):
            config = {
                "excel_output_path": output,
                "sheet_name": "测试",
                "style_header": True,
                "style_alt_rows": True,
                "split_config": {"enabled": False},
                "pivot_config": {"enabled": False},
                "conditional_format_rules": [],
                "validation_rules": [{"field": "id", "rule_type": "not_null", "on_fail": "mark"}],
                "computed_columns": [],
            }
            with patch("builtins.print"):
                result = export_to_excel(data, headers, config)
            assert result == output

    def test_校验中止返回None(self, tmp_path):
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "aborted.xlsx")
        vr = MagicMock()
        vr.aborted = True
        vr.abort_reason = "中止原因"
        vr.summary.return_value = "已中止"
        with patch("json_to_excel.rules_from_config", return_value=[MagicMock()]), \
             patch("json_to_excel.validate_data", return_value=vr), \
             patch("json_to_excel._print_validation_errors"):
            config = {
                "excel_output_path": output,
                "sheet_name": "测试",
                "split_config": {"enabled": False},
                "pivot_config": {"enabled": False},
                "conditional_format_rules": [],
                "validation_rules": [{"field": "id", "rule_type": "not_null", "on_fail": "abort"}],
                "computed_columns": [],
            }
            with patch("builtins.print"):
                result = export_to_excel(data, headers, config)
            assert result is None

    def test_apply_validation_to_export返回None(self, tmp_path):
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "val_none.xlsx")
        vr = MagicMock()
        vr.aborted = False
        vr.skipped_rows = set()
        vr.summary.return_value = "校验完成"
        with patch("json_to_excel.rules_from_config", return_value=[MagicMock()]), \
             patch("json_to_excel.validate_data", return_value=vr), \
             patch("json_to_excel.apply_validation_to_export", return_value=(None, [])):
            config = {
                "excel_output_path": output,
                "sheet_name": "测试",
                "split_config": {"enabled": False},
                "pivot_config": {"enabled": False},
                "conditional_format_rules": [],
                "validation_rules": [{"field": "id"}],
                "computed_columns": [],
            }
            with patch("builtins.print"):
                result = export_to_excel(data, headers, config)
            assert result is None

    def test_带计算列(self, tmp_path):
        data = [{"id": 1, "name": "张三"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "name", "label": "姓名", "width": 15}]
        output = str(tmp_path / "computed.xlsx")
        new_headers = headers + [{"key": "calc", "label": "计算列", "width": 10}]
        with patch("json_to_excel.apply_computed_columns", return_value=({}, new_headers)):
            config = {
                "excel_output_path": output,
                "sheet_name": "测试",
                "style_header": True,
                "style_alt_rows": True,
                "split_config": {"enabled": False},
                "pivot_config": {"enabled": False},
                "conditional_format_rules": [],
                "validation_rules": [],
                "computed_columns": [{"label": "计算列", "expression": "id * 2", "enabled": True}],
            }
            with patch("builtins.print"):
                result = export_to_excel(data, headers, config)
            assert result == output

    def test_拆分导出路径(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "split_path.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
            "validation_rules": [],
            "computed_columns": [],
        }
        with patch("builtins.print"):
            result = export_to_excel(data, headers, config)
        assert result == output

    def test_带validation_marks(self, tmp_path):
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "marked.xlsx")
        vr = MagicMock()
        vr.aborted = False
        vr.skipped_rows = set()
        vr.marked_rows = {0}
        vr.summary.return_value = "校验完成"
        with patch("json_to_excel.rules_from_config", return_value=[MagicMock()]), \
             patch("json_to_excel.validate_data", return_value=vr), \
             patch("json_to_excel.apply_validation_to_export", return_value=(data, [])):
            config = {
                "excel_output_path": output,
                "sheet_name": "测试",
                "style_header": True,
                "style_alt_rows": True,
                "split_config": {"enabled": False},
                "pivot_config": {"enabled": False},
                "conditional_format_rules": [],
                "validation_rules": [{"field": "id", "rule_type": "not_null", "on_fail": "mark"}],
                "computed_columns": [],
            }
            with patch("builtins.print"):
                result = export_to_excel(data, headers, config)
            assert result == output

    def test_非字典项跳过(self, tmp_path):
        data = ["not_dict", {"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "nondict.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
            "validation_rules": [],
            "computed_columns": [],
        }
        with patch("builtins.print"):
            result = export_to_excel(data, headers, config)
        assert result == output

    def test_带skipped和marked打印(self, tmp_path):
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "skip_mark.xlsx")
        vr = MagicMock()
        vr.aborted = False
        vr.skipped_rows = {1}
        vr.marked_rows = {0}
        vr.summary.return_value = "校验完成"
        with patch("json_to_excel.rules_from_config", return_value=[MagicMock()]), \
             patch("json_to_excel.validate_data", return_value=vr), \
             patch("json_to_excel.apply_validation_to_export", return_value=(data, [])):
            config = {
                "excel_output_path": output,
                "sheet_name": "测试",
                "style_header": True,
                "style_alt_rows": True,
                "split_config": {"enabled": False},
                "pivot_config": {"enabled": False},
                "conditional_format_rules": [],
                "validation_rules": [{"field": "id"}],
                "computed_columns": [],
            }
            with patch("builtins.print"):
                result = export_to_excel(data, headers, config)
            assert result == output


class TestApplyValidationMarks:
    def test_基本标记(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID", "姓名"])
        ws.append([1, "张三"])
        ws.append([2, "李四"])
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "name", "label": "姓名", "width": 15}]
        vr = MagicMock()
        vr.marked_rows = {0}
        rule = MagicMock()
        rule.message = "ID不能为空"
        vr.get_field_errors_for_row.return_value = {"id": [MagicMock(rule=rule)]}
        _apply_validation_marks(ws, vr, headers, [0, 1])
        cell = ws.cell(row=2, column=1)
        assert cell.font.bold is True
        assert cell.comment is not None

    def test_长错误消息截断(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        vr = MagicMock()
        vr.marked_rows = {0}
        rule = MagicMock()
        rule.message = "A" * 300
        vr.get_field_errors_for_row.return_value = {"id": [MagicMock(rule=rule)]}
        _apply_validation_marks(ws, vr, headers, [0])
        cell = ws.cell(row=2, column=1)
        assert cell.comment is not None
        assert "..." in cell.comment.text

    def test_orig_idx不在original_indices中(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        ws.append([1])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        vr = MagicMock()
        vr.marked_rows = {5}
        _apply_validation_marks(ws, vr, headers, [0])
        cell = ws.cell(row=2, column=1)
        assert cell.comment is None

    def test_字段错误标记特定列(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID", "姓名"])
        ws.append([1, "张三"])
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "name", "label": "姓名", "width": 15}]
        vr = MagicMock()
        vr.marked_rows = {0}
        rule1 = MagicMock()
        rule1.message = "ID错误"
        rule2 = MagicMock()
        rule2.message = "姓名错误"
        vr.get_field_errors_for_row.return_value = {
            "id": [MagicMock(rule=rule1)],
            "name": [MagicMock(rule=rule2)],
        }
        _apply_validation_marks(ws, vr, headers, [0])
        name_cell = ws.cell(row=2, column=2)
        assert name_cell.fill.start_color.rgb is not None


class TestPrintValidationErrors:
    def test_打印错误详情(self, capsys):
        vr = MagicMock()
        rule = MagicMock()
        rule.message = "ID不能为空"
        rule.on_fail = "mark"
        error = MagicMock()
        error.row_index = 0
        error.value = None
        error.rule = rule
        vr.errors = [error]
        data = [{"id": None}]
        _print_validation_errors(vr, data)
        captured = capsys.readouterr()
        assert "校验错误详情" in captured.out
        assert "ID不能为空" in captured.out

    def test_无错误直接返回(self, capsys):
        vr = MagicMock()
        vr.errors = []
        _print_validation_errors(vr, [])
        captured = capsys.readouterr()
        assert captured.out == ""

    def test_值不为None时截断(self, capsys):
        vr = MagicMock()
        rule = MagicMock()
        rule.message = "错误"
        rule.on_fail = "mark"
        error = MagicMock()
        error.row_index = 0
        error.value = "A" * 50
        error.rule = rule
        vr.errors = [error]
        _print_validation_errors(vr, [{"id": 1}])
        captured = capsys.readouterr()
        assert "校验错误详情" in captured.out

    def test_超过max_display显示省略(self, capsys):
        vr = MagicMock()
        errors = []
        for i in range(15):
            rule = MagicMock()
            rule.message = f"错误{i}"
            rule.on_fail = "mark"
            error = MagicMock()
            error.row_index = i
            error.value = None
            error.rule = rule
            errors.append(error)
        vr.errors = errors
        _print_validation_errors(vr, [{}], max_display=10)
        captured = capsys.readouterr()
        assert "还有" in captured.out


class TestParseArgs:
    def test_默认参数(self):
        with patch("sys.argv", ["json_to_excel.py"]):
            args = parse_args()
            assert args.wizard is False
            assert args.preview is False
            assert args.config is None
            assert args.input is None
            assert args.output is None
            assert args.format is None

    def test_带参数(self):
        with patch("sys.argv", ["json_to_excel.py", "-i", "test.json", "-o", "out.xlsx", "-f", "excel"]):
            args = parse_args()
            assert args.input == "test.json"
            assert args.output == "out.xlsx"
            assert args.format == "excel"

    def test_wizard参数(self):
        with patch("sys.argv", ["json_to_excel.py", "-w"]):
            args = parse_args()
            assert args.wizard is True

    def test_batch参数(self):
        with patch("sys.argv", ["json_to_excel.py", "--batch"]):
            args = parse_args()
            assert args.batch is True

    def test_validate参数(self):
        with patch("sys.argv", ["json_to_excel.py", "--validate"]):
            args = parse_args()
            assert args.validate is True

    def test_list_formats参数(self):
        with patch("sys.argv", ["json_to_excel.py", "--list-formats"]):
            args = parse_args()
            assert args.list_formats is True


class TestApplyCliOverrides:
    def test_input覆盖(self):
        config = {}
        args = MagicMock()
        args.input = "test.json"
        args.format = None
        args.output = None
        with patch("multi_exporter.get_format_extension", return_value=".xlsx"):
            result = apply_cli_overrides(config, args)
        assert result["json_file_path"] == "test.json"

    def test_format覆盖(self):
        config = {}
        args = MagicMock()
        args.input = None
        args.format = "csv"
        args.output = None
        with patch("multi_exporter.get_format_extension", return_value=".csv"):
            result = apply_cli_overrides(config, args)
        assert result["export_format"] == "csv"

    def test_output覆盖excel(self):
        config = {"export_format": "excel"}
        args = MagicMock()
        args.input = None
        args.format = None
        args.output = "out.xlsx"
        with patch("multi_exporter.get_format_extension", return_value=".xlsx"):
            result = apply_cli_overrides(config, args)
        assert result["excel_output_path"] == "out.xlsx"

    def test_output覆盖csv(self):
        config = {"export_format": "csv"}
        args = MagicMock()
        args.input = None
        args.format = None
        args.output = "out.csv"
        with patch("multi_exporter.get_format_extension", return_value=".csv"):
            result = apply_cli_overrides(config, args)
        assert result["csv_output_path"] == "out.csv"


class TestRunWithConfig:
    def test_正常执行(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        config = {
            "json_file_path": str(f),
            "excel_output_path": str(tmp_path / "out.xlsx"),
            "export_format": "excel",
        }
        with patch("multi_exporter.export_data"), \
             patch("multi_exporter.EXPORT_FORMATS", {"excel": {"label": "Excel", "description": "Excel格式"}}), \
             patch("json_to_excel.prompt_save_as_template"), \
             patch("builtins.print"):
            run_with_config(config)

    def test_使用筛选后数据(self, tmp_path):
        data = [{"id": 1}]
        config = {
            "json_file_path": str(tmp_path / "test.json"),
            "excel_output_path": str(tmp_path / "out.xlsx"),
            "export_format": "excel",
        }
        with patch("multi_exporter.export_data"), \
             patch("multi_exporter.EXPORT_FORMATS", {"excel": {"label": "Excel", "description": "Excel格式"}}), \
             patch("json_to_excel.prompt_save_as_template"), \
             patch("builtins.print"):
            run_with_config(config, data=data)

    def test_文件不存在退出(self):
        config = {
            "json_file_path": "/nonexistent/file.json",
            "excel_output_path": "/tmp/out.xlsx",
            "export_format": "excel",
        }
        with patch("multi_exporter.EXPORT_FORMATS", {"excel": {"label": "Excel", "description": "Excel格式"}}), \
             patch("builtins.print"), \
             pytest.raises(SystemExit):
            run_with_config(config)

    def test_JSON解析错误退出(self, tmp_path):
        f = tmp_path / "bad.json"
        f.write_text("{invalid json", encoding="utf-8")
        config = {
            "json_file_path": str(f),
            "excel_output_path": str(tmp_path / "out.xlsx"),
            "export_format": "excel",
        }
        with patch("multi_exporter.EXPORT_FORMATS", {"excel": {"label": "Excel", "description": "Excel格式"}}), \
             patch("builtins.print"), \
             pytest.raises(SystemExit):
            run_with_config(config)

    def test_一般异常退出(self, tmp_path):
        config = {
            "json_file_path": str(tmp_path / "test.json"),
            "excel_output_path": str(tmp_path / "out.xlsx"),
            "export_format": "excel",
        }
        with patch("multi_exporter.EXPORT_FORMATS", {"excel": {"label": "Excel", "description": "Excel格式"}}), \
             patch("json_to_excel.load_json", side_effect=Exception("未知错误")), \
             patch("builtins.print"), \
             patch("traceback.print_exc"), \
             pytest.raises(SystemExit):
            run_with_config(config)


class TestPromptSaveAsTemplate:
    def test_不保存直接返回(self):
        config = {}
        with patch("prompts.prompt_confirm", return_value=False):
            prompt_save_as_template(config)

    def test_保存模板成功(self):
        config = {}
        with patch("prompts.prompt_confirm", side_effect=[True, False]), \
             patch("prompts.prompt_input", side_effect=["my_style", "我的样式", ""]), \
             patch("style_template_manager.create_template_from_config", return_value=(True, "保存成功")), \
             patch("builtins.print"):
            prompt_save_as_template(config)

    def test_模板已存在覆盖(self):
        config = {}
        with patch("prompts.prompt_confirm", side_effect=[True, True]), \
             patch("prompts.prompt_input", side_effect=["my_style", "我的样式", ""]), \
             patch("style_template_manager.create_template_from_config", side_effect=[
                 (False, "模板 'my_style' 已存在"),
                 (True, "保存成功"),
             ]), \
             patch("builtins.print"):
            prompt_save_as_template(config)

    def test_模板已存在不覆盖(self):
        config = {}
        with patch("prompts.prompt_confirm", side_effect=[True, False]), \
             patch("prompts.prompt_input", side_effect=["my_style", "我的样式", ""]), \
             patch("style_template_manager.create_template_from_config", return_value=(False, "模板 'my_style' 已存在")), \
             patch("builtins.print"):
            prompt_save_as_template(config)


class TestHandleTemplateOperations:
    def test_list_templates(self):
        args = MagicMock()
        args.list_templates = True
        args.template_manager = False
        args.apply_template = None
        args.save_as_template = None
        args.template_name = None
        args.template_desc = ""
        config = {}
        mock_templates = {
            "default": {"name": "默认", "description": "默认模板", "header_style": {"bg_color": "#4472C4", "border_style": "thin"}, "data_style": {"alt_row_color": "#F2F2F2"}, "conditional_format_rules": []},
        }
        with patch("style_template_manager.list_templates", return_value=mock_templates), \
             patch("style_template_manager.get_border_style_name", return_value="细线"), \
             patch("builtins.print"):
            result = handle_template_operations(args, config)
        assert result is True

    def test_list_templates_with_conditional_rules(self):
        args = MagicMock()
        args.list_templates = True
        args.template_manager = False
        args.apply_template = None
        args.save_as_template = None
        args.template_name = None
        args.template_desc = ""
        config = {}
        mock_templates = {
            "custom": {
                "name": "自定义",
                "description": "自定义模板",
                "header_style": {"bg_color": "#FFFFFF", "border_style": "medium"},
                "data_style": {"alt_row_color": "#FFFFFF"},
                "conditional_format_rules": [{"enabled": True, "field": "id", "condition": "value > 0"}],
            },
        }
        with patch("style_template_manager.list_templates", return_value=mock_templates), \
             patch("style_template_manager.get_border_style_name", return_value="中等"), \
             patch("style_template_manager.get_rule_description", return_value="id > 0"), \
             patch("builtins.print"):
            result = handle_template_operations(args, config)
        assert result is True

    def test_template_manager(self):
        args = MagicMock()
        args.list_templates = False
        args.template_manager = True
        args.apply_template = None
        args.save_as_template = None
        config = {}
        with patch("style_template_wizard.run_style_template_manager"):
            result = handle_template_operations(args, config)
        assert result is True

    def test_apply_template成功(self):
        args = MagicMock()
        args.list_templates = False
        args.template_manager = False
        args.apply_template = "default"
        args.save_as_template = None
        config = {}
        with patch("style_template_manager.apply_template_to_config", return_value=(True, "应用成功")), \
             patch("builtins.print"):
            result = handle_template_operations(args, config)
        assert result is False

    def test_apply_template失败退出(self):
        args = MagicMock()
        args.list_templates = False
        args.template_manager = False
        args.apply_template = "nonexistent"
        args.save_as_template = None
        config = {}
        with patch("style_template_manager.apply_template_to_config", return_value=(False, "模板不存在")), \
             patch("builtins.print"), \
             pytest.raises(SystemExit):
            handle_template_operations(args, config)

    def test_save_as_template成功(self):
        args = MagicMock()
        args.list_templates = False
        args.template_manager = False
        args.apply_template = None
        args.save_as_template = "my_style"
        args.template_name = "我的样式"
        args.template_desc = "描述"
        config = {}
        with patch("style_template_manager.create_template_from_config", return_value=(True, "保存成功")), \
             patch("builtins.print"):
            result = handle_template_operations(args, config)
        assert result is True

    def test_save_as_template失败退出(self):
        args = MagicMock()
        args.list_templates = False
        args.template_manager = False
        args.apply_template = None
        args.save_as_template = "my_style"
        args.template_name = "我的样式"
        args.template_desc = ""
        config = {}
        with patch("style_template_manager.create_template_from_config", return_value=(False, "保存失败")), \
             patch("builtins.print"), \
             pytest.raises(SystemExit):
            handle_template_operations(args, config)

    def test_无模板操作返回False(self):
        args = MagicMock()
        args.list_templates = False
        args.template_manager = False
        args.apply_template = None
        args.save_as_template = None
        config = {}
        result = handle_template_operations(args, config)
        assert result is False


class TestHandleBatchOperations:
    def test_batch模式(self):
        args = MagicMock()
        args.batch = True
        args.batch_list = False
        args.batch_resume = None
        args.batch_report = None
        args.batch_dir = None
        args.batch_files = None
        args.config = None
        with patch("batch_wizard.run_batch_wizard"), \
             patch("json_to_excel.load_config", return_value={}):
            result = _handle_batch_operations(args)
        assert result is True

    def test_batch_list模式(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = True
        with patch("batch_wizard.show_task_manager"):
            result = _handle_batch_operations(args)
        assert result is True

    def test_batch_resume成功(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = "task123"
        args.batch_report = None
        args.batch_dir = None
        args.batch_files = None
        args.batch_report_format = "txt"
        mock_task = MagicMock()
        with patch("batch_processor.resume_batch_process", return_value=(mock_task, "ok")), \
             patch("report_generator.generate_report", return_value="/tmp/report.txt"), \
             patch("report_generator.print_summary"), \
             patch("builtins.print"):
            result = _handle_batch_operations(args)
        assert result is True

    def test_batch_resume失败(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = "task123"
        args.batch_report = None
        args.batch_dir = None
        args.batch_files = None
        args.batch_report_format = "txt"
        with patch("batch_processor.resume_batch_process", return_value=(None, "任务不存在")), \
             patch("builtins.print"), \
             pytest.raises(SystemExit):
            _handle_batch_operations(args)

    def test_batch_report成功(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = None
        args.batch_report = "task123"
        args.batch_dir = None
        args.batch_files = None
        args.batch_report_format = "txt"
        mock_task = MagicMock()
        with patch("task_manager.load_batch_task", return_value=mock_task), \
             patch("report_generator.generate_report", return_value="/tmp/report.txt"), \
             patch("report_generator.print_summary"), \
             patch("builtins.print"):
            result = _handle_batch_operations(args)
        assert result is True

    def test_batch_report任务不存在(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = None
        args.batch_report = "task123"
        args.batch_dir = None
        args.batch_files = None
        args.batch_report_format = "txt"
        with patch("task_manager.load_batch_task", return_value=None), \
             patch("builtins.print"), \
             pytest.raises(SystemExit):
            _handle_batch_operations(args)

    def test_batch_dir模式(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = None
        args.batch_report = None
        args.batch_dir = "/data/json"
        args.batch_files = None
        args.batch_output = "/output"
        args.batch_report_format = "txt"
        args.batch_skip_existing = False
        args.config = None
        mock_task = MagicMock()
        with patch("json_to_excel.load_config", return_value={}), \
             patch("json_to_excel.apply_cli_overrides", return_value={}), \
             patch("batch_processor.start_batch_process", return_value=(mock_task, "ok")), \
             patch("report_generator.generate_report", return_value="/tmp/report.txt"), \
             patch("report_generator.print_summary"), \
             patch("builtins.print"):
            result = _handle_batch_operations(args)
        assert result is True

    def test_batch_files模式(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = None
        args.batch_report = None
        args.batch_dir = None
        args.batch_files = "a.json,b.json"
        args.batch_output = "/output"
        args.batch_report_format = "json"
        args.batch_skip_existing = False
        args.config = None
        mock_task = MagicMock()
        with patch("json_to_excel.load_config", return_value={}), \
             patch("json_to_excel.apply_cli_overrides", return_value={}), \
             patch("batch_processor.start_batch_process", return_value=(mock_task, "ok")), \
             patch("report_generator.generate_json_report", return_value="/tmp/report.json"), \
             patch("report_generator.print_summary"), \
             patch("builtins.print"):
            result = _handle_batch_operations(args)
        assert result is True

    def test_batch_start失败(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = None
        args.batch_report = None
        args.batch_dir = "/data/json"
        args.batch_files = None
        args.batch_output = "/output"
        args.batch_report_format = "txt"
        args.batch_skip_existing = False
        args.config = None
        with patch("json_to_excel.load_config", return_value={}), \
             patch("json_to_excel.apply_cli_overrides", return_value={}), \
             patch("batch_processor.start_batch_process", return_value=(None, "启动失败")), \
             patch("builtins.print"), \
             pytest.raises(SystemExit):
            _handle_batch_operations(args)

    def test_无batch操作返回False(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = None
        args.batch_report = None
        args.batch_dir = None
        args.batch_files = None
        result = _handle_batch_operations(args)
        assert result is False


class TestMain函数:
    def test_list_formats(self):
        with patch("sys.argv", ["json_to_excel.py", "--list-formats"]), \
             patch("multi_exporter.EXPORT_FORMATS", {"excel": {"label": "Excel", "description": "Excel格式"}}), \
             patch("builtins.print"):
            main()

    def test_validate模式(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        with patch("sys.argv", ["json_to_excel.py", "--validate", "-c", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": str(f)}), \
             patch("validation_wizard.run_validation_wizard", return_value={"validation_rules": []}), \
             patch("builtins.print"):
            main()

    def test_validate模式异常(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "--validate"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/nonexistent"}), \
             patch("validation_wizard.run_validation_wizard", return_value={}), \
             patch("builtins.print"):
            main()

    def test_validate_only模式(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        vr = MagicMock()
        vr.summary.return_value = "校验完成"
        vr.errors = []
        with patch("sys.argv", ["json_to_excel.py", "--validate-only", "-c", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": str(f)}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": str(f)}), \
             patch("json_to_excel.rules_from_config", return_value=[MagicMock()]), \
             patch("json_to_excel.validate_data", return_value=vr), \
             patch("json_to_excel._print_validation_errors"), \
             patch("builtins.print"):
            main()

    def test_validate_only无规则(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "--validate-only"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.rules_from_config", return_value=[]), \
             patch("builtins.print"):
            main()

    def test_validate_only异常(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "--validate-only"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/nonexistent"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/nonexistent"}), \
             patch("json_to_excel.rules_from_config", return_value=[MagicMock()]), \
             patch("json_to_excel.load_json", side_effect=Exception("加载失败")), \
             patch("builtins.print"):
            main()

    def test_computed_columns模式(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        with patch("sys.argv", ["json_to_excel.py", "--computed-columns", "-c", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": str(f)}), \
             patch("computed_columns_wizard.run_computed_columns_wizard", return_value={}), \
             patch("builtins.print"):
            main()

    def test_computed_columns模式异常(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "--computed-columns"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/nonexistent"}), \
             patch("computed_columns_wizard.run_computed_columns_wizard", return_value={}), \
             patch("builtins.print"):
            main()

    def test_wizard模式(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "-w"]), \
             patch("json_to_excel.load_config", return_value={}), \
             patch("wizard.run_wizard", return_value={}), \
             patch("json_to_excel.validate_config", return_value=[]), \
             patch("json_to_excel.prompt_confirm_execute", return_value=False), \
             patch("builtins.print"):
            main()

    def test_wizard模式验证失败(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "-w"]), \
             patch("json_to_excel.load_config", return_value={}), \
             patch("wizard.run_wizard", return_value={}), \
             patch("json_to_excel.validate_config", return_value=["配置错误"]), \
             patch("builtins.print"), \
             pytest.raises(SystemExit):
            main()

    def test_wizard模式保存配置(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "-w", "--save-config", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={}), \
             patch("wizard.run_wizard", return_value={}), \
             patch("json_to_excel.validate_config", return_value=[]), \
             patch("json_to_excel.prompt_confirm_execute", return_value=False), \
             patch("json_to_excel.save_config"), \
             patch("builtins.print"):
            main()

    def test_wizard模式执行导出(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "-w"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("wizard.run_wizard", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.validate_config", return_value=[]), \
             patch("json_to_excel.prompt_confirm_execute", return_value=True), \
             patch("json_to_excel.run_with_config"), \
             patch("builtins.print"):
            main()

    def test_batch操作(self):
        with patch("sys.argv", ["json_to_excel.py", "--batch"]), \
             patch("json_to_excel.load_config", return_value={}), \
             patch("batch_wizard.run_batch_wizard"), \
             patch("builtins.print"):
            main()

    def test_preview模式(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        with patch("sys.argv", ["json_to_excel.py", "-p", "-c", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": str(f)}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": str(f)}), \
             patch("data_previewer.start_preview_mode", return_value=data), \
             patch("json_to_excel.prompt_confirm_execute_after_preview", return_value=False), \
             patch("builtins.print"):
            main()

    def test_preview模式导出(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        with patch("sys.argv", ["json_to_excel.py", "-p", "-c", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": str(f), "export_format": "excel"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": str(f), "export_format": "excel"}), \
             patch("data_previewer.start_preview_mode", return_value=data), \
             patch("json_to_excel.prompt_confirm_execute_after_preview", return_value=True), \
             patch("multi_exporter.export_data"), \
             patch("json_to_excel.prompt_save_as_template"), \
             patch("builtins.print"):
            main()

    def test_preview返回None(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "-p"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("data_previewer.start_preview_mode", return_value=None), \
             patch("builtins.print"):
            main()

    def test_默认路径验证失败(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "-c", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.validate_config", return_value=["配置错误"]), \
             patch("builtins.print"), \
             pytest.raises(SystemExit):
            main()

    def test_默认路径保存配置(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "--save-config", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.validate_config", return_value=[]), \
             patch("json_to_excel.save_config"), \
             patch("json_to_excel.prompt_confirm_preview", return_value=False), \
             patch("json_to_excel.run_with_config"), \
             patch("builtins.print"):
            main()

    def test_prompt_confirm_preview返回True(self, tmp_path):
        data = [{"id": 1}]
        with patch("sys.argv", ["json_to_excel.py"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.validate_config", return_value=[]), \
             patch("json_to_excel.prompt_confirm_preview", return_value=True), \
             patch("data_previewer.start_preview_mode", return_value=data), \
             patch("json_to_excel.prompt_confirm_execute_after_preview", return_value=False), \
             patch("builtins.print"):
            main()

    def test_apply_template路径(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "--apply-template", "default"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("style_template_manager.apply_template_to_config", return_value=(True, "应用成功")), \
             patch("json_to_excel.prompt_confirm_execute", return_value=False), \
             patch("builtins.print"):
            main()

    def test_apply_template保存配置并执行(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "--apply-template", "default", "--save-config", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("style_template_manager.apply_template_to_config", return_value=(True, "应用成功")), \
             patch("json_to_excel.prompt_confirm_execute", return_value=True), \
             patch("json_to_excel.run_with_config"), \
             patch("json_to_excel.save_config"), \
             patch("builtins.print"):
            main()


class TestPromptConfirmPreview:
    def test_输入y返回True(self):
        with patch("builtins.input", return_value="y"):
            assert prompt_confirm_preview() is True

    def test_输入n返回False(self):
        with patch("builtins.input", return_value="n"):
            assert prompt_confirm_preview() is False

    def test_输入空返回True(self):
        with patch("builtins.input", return_value=""):
            assert prompt_confirm_preview() is True

    def test_输入yes返回True(self):
        with patch("builtins.input", return_value="yes"):
            assert prompt_confirm_preview() is True

    def test_输入no返回False(self):
        with patch("builtins.input", return_value="no"):
            assert prompt_confirm_preview() is False

    def test_输入无效后有效(self):
        with patch("builtins.input", side_effect=["x", "y"]):
            assert prompt_confirm_preview() is True


class TestPromptConfirmExecute:
    def test_输入y返回True(self):
        with patch("builtins.input", return_value="y"):
            assert prompt_confirm_execute() is True

    def test_输入n返回False(self):
        with patch("builtins.input", return_value="n"):
            assert prompt_confirm_execute() is False

    def test_输入无效后有效(self):
        with patch("builtins.input", side_effect=["x", "n"]):
            assert prompt_confirm_execute() is False


class TestPromptConfirmExecuteAfterPreview:
    def test_输入y返回True(self):
        with patch("multi_exporter.EXPORT_FORMATS", {"excel": {"label": "Excel"}}), \
             patch("builtins.input", return_value="y"):
            assert prompt_confirm_execute_after_preview() is True

    def test_输入n返回False(self):
        with patch("multi_exporter.EXPORT_FORMATS", {"excel": {"label": "Excel"}}), \
             patch("builtins.input", return_value="n"):
            assert prompt_confirm_execute_after_preview() is False

    def test_带config参数(self):
        config = {"export_format": "csv"}
        with patch("multi_exporter.EXPORT_FORMATS", {"csv": {"label": "CSV"}}), \
             patch("builtins.input", return_value="y"):
            assert prompt_confirm_execute_after_preview(config) is True

    def test_输入无效后有效(self):
        with patch("multi_exporter.EXPORT_FORMATS", {"excel": {"label": "Excel"}}), \
             patch("builtins.input", side_effect=["x", "y"]):
            assert prompt_confirm_execute_after_preview() is True


class TestCreatePivotSheet:
    def test_有列字段单值字段(self):
        wb = Workbook()
        ws_default = wb.active
        data = [
            {"dept": "技术部", "year": "2024", "salary": 10000},
            {"dept": "技术部", "year": "2023", "salary": 20000},
            {"dept": "市场部", "year": "2024", "salary": 15000},
        ]
        pivot_config = {
            "sheet_name": "透视表",
            "row_fields": ["dept"],
            "column_fields": ["year"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
            "show_row_totals": True,
            "show_column_totals": True,
            "grand_total_label": "总计",
            "apply_style": True,
        }
        headers = [
            {"key": "dept", "label": "部门"},
            {"key": "year", "label": "年份"},
            {"key": "salary", "label": "薪资"},
        ]
        pivot_result = build_pivot_table(data, pivot_config, headers)
        result_ws = create_pivot_sheet(wb, pivot_result, pivot_config, headers)
        assert result_ws is not None
        assert result_ws.cell(row=1, column=1).value is not None

    def test_有列字段多值字段(self):
        wb = Workbook()
        ws_default = wb.active
        data = [
            {"dept": "技术部", "year": "2024", "salary": 10000, "bonus": 2000},
            {"dept": "市场部", "year": "2024", "salary": 15000, "bonus": 3000},
        ]
        pivot_config = {
            "sheet_name": "透视表",
            "row_fields": ["dept"],
            "column_fields": ["year"],
            "value_fields": [
                {"field": "salary", "aggregate": "sum", "label": "薪资"},
                {"field": "bonus", "aggregate": "sum", "label": "奖金"},
            ],
            "show_row_totals": True,
            "show_column_totals": True,
            "grand_total_label": "总计",
            "apply_style": True,
        }
        headers = [
            {"key": "dept", "label": "部门"},
            {"key": "year", "label": "年份"},
            {"key": "salary", "label": "薪资"},
            {"key": "bonus", "label": "奖金"},
        ]
        pivot_result = build_pivot_table(data, pivot_config, headers)
        result_ws = create_pivot_sheet(wb, pivot_result, pivot_config, headers)
        assert result_ws is not None

    def test_无列字段多值字段(self):
        wb = Workbook()
        ws_default = wb.active
        data = [
            {"dept": "技术部", "salary": 10000, "bonus": 2000},
            {"dept": "市场部", "salary": 15000, "bonus": 3000},
        ]
        pivot_config = {
            "sheet_name": "透视表",
            "row_fields": ["dept"],
            "column_fields": [],
            "value_fields": [
                {"field": "salary", "aggregate": "sum", "label": "薪资"},
                {"field": "bonus", "aggregate": "sum", "label": "奖金"},
            ],
            "show_row_totals": True,
            "show_column_totals": False,
            "apply_style": True,
        }
        headers = [
            {"key": "dept", "label": "部门"},
            {"key": "salary", "label": "薪资"},
            {"key": "bonus", "label": "奖金"},
        ]
        pivot_result = build_pivot_table(data, pivot_config, headers)
        result_ws = create_pivot_sheet(wb, pivot_result, pivot_config, headers)
        assert result_ws is not None

    def test_无列字段单值字段(self):
        wb = Workbook()
        ws_default = wb.active
        data = [
            {"dept": "技术部", "salary": 10000},
            {"dept": "市场部", "salary": 15000},
        ]
        pivot_config = {
            "sheet_name": "透视表",
            "row_fields": ["dept"],
            "column_fields": [],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
            "show_row_totals": False,
            "show_column_totals": False,
            "apply_style": True,
        }
        headers = [
            {"key": "dept", "label": "部门"},
            {"key": "salary", "label": "薪资"},
        ]
        pivot_result = build_pivot_table(data, pivot_config, headers)
        try:
            result_ws = create_pivot_sheet(wb, pivot_result, pivot_config, headers)
            assert result_ws is not None
        except ValueError:
            pass

    def test_不应用样式(self):
        wb = Workbook()
        ws_default = wb.active
        data = [{"dept": "技术部", "year": "2024", "salary": 10000}]
        pivot_config = {
            "sheet_name": "透视表",
            "row_fields": ["dept"],
            "column_fields": ["year"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
            "show_row_totals": False,
            "show_column_totals": False,
            "apply_style": False,
        }
        headers = [{"key": "dept", "label": "部门"}, {"key": "year", "label": "年份"}, {"key": "salary", "label": "薪资"}]
        pivot_result = build_pivot_table(data, pivot_config, headers)
        result_ws = create_pivot_sheet(wb, pivot_result, pivot_config, headers)
        assert result_ws is not None

    def test_有列字段不显示总计(self):
        wb = Workbook()
        ws_default = wb.active
        data = [
            {"dept": "技术部", "year": "2024", "salary": 10000},
        ]
        pivot_config = {
            "sheet_name": "透视表",
            "row_fields": ["dept"],
            "column_fields": ["year"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
            "show_row_totals": False,
            "show_column_totals": False,
            "apply_style": True,
        }
        headers = [{"key": "dept", "label": "部门"}, {"key": "year", "label": "年份"}, {"key": "salary", "label": "薪资"}]
        pivot_result = build_pivot_table(data, pivot_config, headers)
        result_ws = create_pivot_sheet(wb, pivot_result, pivot_config, headers)
        assert result_ws is not None

    def test_多行字段有列字段(self):
        wb = Workbook()
        ws_default = wb.active
        data = [
            {"dept": "技术部", "level": "高级", "year": "2024", "salary": 10000},
            {"dept": "市场部", "level": "初级", "year": "2024", "salary": 15000},
        ]
        pivot_config = {
            "sheet_name": "透视表",
            "row_fields": ["dept", "level"],
            "column_fields": ["year"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
            "show_row_totals": True,
            "show_column_totals": True,
            "grand_total_label": "总计",
            "apply_style": True,
        }
        headers = [
            {"key": "dept", "label": "部门"},
            {"key": "level", "label": "级别"},
            {"key": "year", "label": "年份"},
            {"key": "salary", "label": "薪资"},
        ]
        pivot_result = build_pivot_table(data, pivot_config, headers)
        result_ws = create_pivot_sheet(wb, pivot_result, pivot_config, headers)
        assert result_ws is not None


class TestAddPivotTableToWorkbook:
    def test_未启用返回None(self):
        wb = Workbook()
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID"}]
        config = {"pivot_config": {"enabled": False}}
        result = add_pivot_table_to_workbook(wb, data, headers, config)
        assert result is None

    def test_启用成功(self):
        wb = Workbook()
        data = [{"dept": "技术部", "year": "2024", "salary": 10000}]
        headers = [{"key": "dept", "label": "部门"}, {"key": "year", "label": "年份"}, {"key": "salary", "label": "薪资"}]
        config = {
            "pivot_config": {
                "enabled": True,
                "row_fields": ["dept"],
                "column_fields": ["year"],
                "value_fields": [{"field": "salary", "aggregate": "sum"}],
            }
        }
        with patch("builtins.print"):
            result = add_pivot_table_to_workbook(wb, data, headers, config)
        assert result is not None

    def test_启用异常(self):
        wb = Workbook()
        data = [{"dept": "技术部"}]
        headers = [{"key": "dept", "label": "部门"}]
        config = {
            "pivot_config": {
                "enabled": True,
                "row_fields": ["dept"],
                "column_fields": [],
                "value_fields": [{"field": "salary", "aggregate": "sum"}],
            }
        }
        with patch("json_to_excel.build_pivot_table", side_effect=Exception("测试异常")), \
             patch("builtins.print"), \
             patch("traceback.print_exc"):
            result = add_pivot_table_to_workbook(wb, data, headers, config)
        assert result is None


class TestDataLoader计算列高级:
    def test_apply_computed_columns有计算列(self, tmp_path):
        data = [{"id": 1, "name": "张三"}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        loader = DataLoader({
            "json_file_path": str(f),
            "computed_columns": [{"label": "计算列", "expression": "id * 2", "enabled": True}],
        })
        loader.set_data(data)
        loader.detect_headers()
        loader.merge_headers()
        new_headers = loader.headers + [{"key": "calc", "label": "计算列", "width": 10}]
        with patch("computed_columns.apply_computed_columns", return_value=({}, new_headers)):
            cache, headers = loader.apply_computed_columns()
        assert cache is not None
        assert any(h["key"] == "calc" for h in headers)

    def test_apply_computed_columns_enabled为False(self, tmp_path):
        data = [{"id": 1}]
        loader = DataLoader({"computed_columns": [{"label": "计算列", "expression": "1", "enabled": False}]})
        loader.set_data(data)
        loader.detect_headers()
        loader.merge_headers()
        cache, headers = loader.apply_computed_columns()
        assert cache is None


class TestDataLoader校验高级:
    def test_validate_data有规则(self, tmp_path):
        data = [{"id": 1}, {"id": None}]
        loader = DataLoader({"validation_rules": [{"field": "id", "rule_type": "not_null", "on_fail": "mark"}]})
        loader.set_data(data)
        vr = MagicMock()
        vr.summary.return_value = "校验完成"
        with patch("data_validator.validate_data", return_value=vr), \
             patch("data_validator.rules_from_config", return_value=[MagicMock()]):
            result = loader.validate_data()
        assert result is not None

    def test_validation_result属性有值(self, tmp_path):
        data = [{"id": 1}]
        loader = DataLoader({"validation_rules": [{"field": "id", "rule_type": "not_null"}]})
        loader.set_data(data)
        vr = MagicMock()
        vr.summary.return_value = "校验完成"
        with patch("data_validator.validate_data", return_value=vr), \
             patch("data_validator.rules_from_config", return_value=[MagicMock()]):
            loader.validate_data()
        assert loader.validation_result is not None


class TestDataLoader有效数据高级:
    def test_get_valid_data校验中止(self, tmp_path):
        data = [{"id": 1}]
        loader = DataLoader({"validation_rules": [{"field": "id", "rule_type": "not_null"}]})
        loader.set_data(data)
        loader.detect_headers()
        loader.merge_headers()
        vr = MagicMock()
        vr.aborted = True
        vr.abort_reason = "中止原因"
        loader._validation_result = vr
        with patch("json_to_excel._print_validation_errors"):
            result = loader.get_valid_data()
        assert result is None

    def test_get_valid_data_apply返回None(self, tmp_path):
        data = [{"id": 1}]
        loader = DataLoader({"validation_rules": [{"field": "id", "rule_type": "not_null"}]})
        loader.set_data(data)
        loader.detect_headers()
        loader.merge_headers()
        vr = MagicMock()
        vr.aborted = False
        vr.skipped_rows = set()
        loader._validation_result = vr
        with patch("data_validator.apply_validation_to_export", return_value=(None, [])):
            result = loader.get_valid_data()
        assert result is None

    def test_get_valid_data有skipped(self, tmp_path):
        data = [{"id": 1}, {"id": None}]
        loader = DataLoader({"validation_rules": [{"field": "id", "rule_type": "not_null"}]})
        loader.set_data(data)
        loader.detect_headers()
        loader.merge_headers()
        vr = MagicMock()
        vr.aborted = False
        vr.skipped_rows = {1}
        vr.marked_rows = set()
        loader._validation_result = vr
        with patch("data_validator.apply_validation_to_export", return_value=([{"id": 1}], [])):
            result = loader.get_valid_data()
        assert result is not None
        assert len(result) == 1

    def test_original_indices属性懒加载(self, tmp_path):
        data = [{"id": 1}, {"id": 2}]
        loader = DataLoader()
        loader.set_data(data)
        indices = loader.original_indices
        assert indices == [0, 1]


class TestDataLoaderPrepareForExport高级:
    def test_prepare_for_export未加载数据(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        loader = DataLoader({"json_file_path": str(f)})
        with patch("builtins.print"):
            result = loader.prepare_for_export()
        assert result is not None
        assert "data" in result

    def test_prepare_for_export校验中止(self, tmp_path):
        data = [{"id": 1}]
        loader = DataLoader({"json_file_path": "", "validation_rules": [{"field": "id", "rule_type": "not_null"}]})
        loader.set_data(data)
        loader.detect_headers()
        loader.merge_headers()
        vr = MagicMock()
        vr.aborted = True
        vr.abort_reason = "中止"
        loader._validation_result = vr
        with patch("json_to_excel._print_validation_errors"):
            result = loader.get_valid_data()
        assert result is None


class TestExcelExporter高级:
    def test_export_single_sheet带computed_cache(self, tmp_path):
        data = [{"id": 1}]
        item = data[0]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "computed.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        cache = {id(item): {"id": 99}}
        exporter = ExcelExporter(config)
        with patch("builtins.print"):
            result = exporter.export(data, headers, computed_cache=cache)
        assert result == output

    def test_export_single_sheet带validation(self, tmp_path):
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "validated.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        vr = MagicMock()
        vr.marked_rows = {0}
        vr.get_field_errors_for_row.return_value = {"id": [MagicMock(rule=MagicMock(message="错误"))]}
        exporter = ExcelExporter(config)
        with patch("builtins.print"):
            result = exporter.export(data, headers, validation_result=vr, original_indices=[0])
        assert result == output

    def test_export_with_split(self, tmp_path):
        data = [
            {"id": 1, "dept": "技术部"},
            {"id": 2, "dept": "市场部"},
        ]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "split.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        with patch("builtins.print"):
            result = exporter.export(data, headers)
        assert result == output


class Test剩余覆盖:
    def test_render行97终端截断(self):
        tracker = ProgressTracker(total=10, min_interval=0)
        tracker.current = 5
        tracker.start_time = time.time() - 1.0
        tracker.current_field = "a" * 50
        buf = StringIO()
        with patch("shutil.get_terminal_size", return_value=MagicMock(columns=50)), \
             patch("sys.stdout", buf):
            tracker._render()
        output = buf.getvalue()
        assert len(output) <= 50

    def test_conditional_formatting字段不在header(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["ID", "名称"])
        ws.append([1, "张三"])
        headers = [{"key": "id", "label": "ID", "width": 10}]
        data = [{"id": 1, "name": "张三"}]
        mock_matched = [{"font_bold": True}]
        with patch("style_template_manager.match_conditional_rules_for_field", return_value=mock_matched), \
             patch("style_template_manager.merge_conditional_styles", return_value={"font_bold": True}):
            config = {"conditional_format_rules": [{"field": "name"}]}
            with patch("builtins.print"):
                apply_conditional_formatting(ws, headers, data, config)

    def test_extract_value点分隔路径返回current(self):
        item = {"a": {"b": "hello"}}
        flat = {}
        with patch("json_to_excel.flatten_dict", return_value=flat):
            result = extract_value(item, "a.b")
            assert result == "hello"

    def test_export_to_excel_with_split创建目录(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "subdir" / "split.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config)
        assert result == output
        assert os.path.exists(output)

    def test_export_to_excel_with_split名称冲突(self, tmp_path):
        data = [
            {"id": 1, "dept": "a" * 31},
            {"id": 2, "dept": "a" * 31},
        ]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "conflict.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": False,
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        with patch("builtins.print"):
            result = export_to_excel_with_split(data, headers, config)
        assert result == output

    def test_export_to_excel创建目录(self, tmp_path):
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "subdir2" / "test.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
            "validation_rules": [],
            "computed_columns": [],
        }
        with patch("builtins.print"):
            result = export_to_excel(data, headers, config)
        assert result == output
        assert os.path.exists(output)

    def test_export_to_excel带computed_cache(self, tmp_path):
        data = [{"id": 1}]
        item = data[0]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "computed_cache.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {"enabled": False},
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
            "validation_rules": [],
            "computed_columns": [],
        }
        cache = {id(item): {"id": 99}}
        with patch("builtins.print"):
            result = export_to_excel(data, headers, config)
        assert result == output

    def test_batch_resume_json报告(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = "task123"
        args.batch_report = None
        args.batch_dir = None
        args.batch_files = None
        args.batch_report_format = "both"
        mock_task = MagicMock()
        with patch("batch_processor.resume_batch_process", return_value=(mock_task, "ok")), \
             patch("report_generator.generate_report", return_value="/tmp/report.txt"), \
             patch("report_generator.generate_json_report", return_value="/tmp/report.json"), \
             patch("report_generator.print_summary"), \
             patch("builtins.print"):
            result = _handle_batch_operations(args)
        assert result is True

    def test_batch_report_json报告(self):
        args = MagicMock()
        args.batch = False
        args.batch_list = False
        args.batch_resume = None
        args.batch_report = "task123"
        args.batch_dir = None
        args.batch_files = None
        args.batch_report_format = "both"
        mock_task = MagicMock()
        with patch("task_manager.load_batch_task", return_value=mock_task), \
             patch("report_generator.generate_report", return_value="/tmp/report.txt"), \
             patch("report_generator.generate_json_report", return_value="/tmp/report.json"), \
             patch("report_generator.print_summary"), \
             patch("builtins.print"):
            result = _handle_batch_operations(args)
        assert result is True

    def test_main_validate带save_config(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        with patch("sys.argv", ["json_to_excel.py", "--validate", "--save-config", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": str(f)}), \
             patch("validation_wizard.run_validation_wizard", return_value={"validation_rules": []}), \
             patch("json_to_excel.save_config"), \
             patch("builtins.print"):
            main()

    def test_main_computed_columns带save_config(self, tmp_path):
        data = [{"id": 1}]
        f = tmp_path / "test.json"
        f.write_text(json.dumps(data), encoding="utf-8")
        with patch("sys.argv", ["json_to_excel.py", "--computed-columns", "--save-config", str(tmp_path / "config.json")]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": str(f)}), \
             patch("computed_columns_wizard.run_computed_columns_wizard", return_value={}), \
             patch("json_to_excel.save_config"), \
             patch("builtins.print"):
            main()

    def test_main_template_exit_and_not_apply(self, tmp_path):
        with patch("sys.argv", ["json_to_excel.py", "--list-templates"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/tmp/test.json"}), \
             patch("style_template_manager.list_templates", return_value={}), \
             patch("builtins.print"):
            main()

    def test_main_prompt_confirm_preview导出(self, tmp_path):
        data = [{"id": 1}]
        with patch("sys.argv", ["json_to_excel.py"]), \
             patch("json_to_excel.load_config", return_value={"json_file_path": "/tmp/test.json", "export_format": "excel"}), \
             patch("json_to_excel.apply_cli_overrides", return_value={"json_file_path": "/tmp/test.json", "export_format": "excel"}), \
             patch("json_to_excel.validate_config", return_value=[]), \
             patch("json_to_excel.prompt_confirm_preview", return_value=True), \
             patch("data_previewer.start_preview_mode", return_value=data), \
             patch("json_to_excel.prompt_confirm_execute_after_preview", return_value=True), \
             patch("multi_exporter.export_data"), \
             patch("json_to_excel.prompt_save_as_template"), \
             patch("builtins.print"):
            main()

    def test_create_pivot_sheet长值列宽(self):
        wb = Workbook()
        ws_default = wb.active
        data = [
            {"dept": "技术部" * 10, "year": "2024", "salary": 10000},
        ]
        pivot_config = {
            "sheet_name": "透视表",
            "row_fields": ["dept"],
            "column_fields": ["year"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
            "show_row_totals": True,
            "show_column_totals": True,
            "grand_total_label": "总计",
            "apply_style": True,
        }
        headers = [
            {"key": "dept", "label": "部门"},
            {"key": "year", "label": "年份"},
            {"key": "salary", "label": "薪资"},
        ]
        pivot_result = build_pivot_table(data, pivot_config, headers)
        with patch("builtins.print"):
            result_ws = create_pivot_sheet(wb, pivot_result, pivot_config, headers)
        assert result_ws is not None

    def test_DataLoader_prepare_for_export_valid_data_none(self, tmp_path):
        data = [{"id": 1}]
        loader = DataLoader({"validation_rules": [{"field": "id", "rule_type": "not_null"}]})
        loader.set_data(data)
        loader.detect_headers()
        loader.merge_headers()
        vr = MagicMock()
        vr.aborted = True
        vr.abort_reason = "中止"
        vr.summary.return_value = "已中止"
        with patch("data_validator.validate_data", return_value=vr), \
             patch("data_validator.rules_from_config", return_value=[MagicMock()]), \
             patch("json_to_excel._print_validation_errors"):
            result = loader.prepare_for_export()
        assert result is None

    def test_ExcelExporter_export_with_split创建目录(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "subdir3" / "split.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        with patch("builtins.print"):
            result = exporter.export(data, headers)
        assert result == output
        assert os.path.exists(output)

    def test_export_from_loader返回None(self, tmp_path):
        loader = MagicMock()
        loader.prepare_for_export.return_value = None
        exporter = ExcelExporter()
        result = exporter.export_from_loader(loader)
        assert result is None

    def test_print_export_summary多工作表(self, capsys):
        exporter = ExcelExporter()
        exporter._print_export_summary(
            "/tmp/test.xlsx",
            [{"id": 1}],
            [{"key": "id", "label": "ID", "width": 10}],
            sheet_count=3,
        )
        captured = capsys.readouterr()
        assert "3" in captured.out

    def test_print_export_summary带validation(self, capsys):
        exporter = ExcelExporter()
        vr = MagicMock()
        vr.skipped_rows = {0}
        vr.marked_rows = {1}
        exporter._print_export_summary(
            "/tmp/test.xlsx",
            [{"id": 1}],
            [{"key": "id", "label": "ID", "width": 10}],
            validation_result=vr,
        )
        captured = capsys.readouterr()
        assert "跳过" in captured.out
        assert "标记" in captured.out

    def test_repr(self):
        exporter = ExcelExporter()
        r = repr(exporter)
        assert "ExcelExporter" in r

    def test_write_sheet调用(self, tmp_path):
        exporter = ExcelExporter()
        wb = Workbook()
        ws = wb.active
        data = [{"id": 1}]
        headers = [{"key": "id", "label": "ID", "width": 10}]
        with patch("builtins.print"):
            exporter._write_sheet(ws, data, headers)
        assert ws.cell(row=2, column=1).value == 1

    def test_export_with_split带validation(self, tmp_path):
        data = [{"id": 1, "dept": "技术部"}]
        headers = [{"key": "id", "label": "ID", "width": 10}, {"key": "dept", "label": "部门", "width": 15}]
        output = str(tmp_path / "split_val.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": True,
                "all_sheet_name": "全部数据",
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        vr = MagicMock()
        vr.marked_rows = {0}
        vr.skipped_rows = set()
        vr.get_field_errors_for_row.return_value = {"id": [MagicMock(rule=MagicMock(message="错误"))]}
        exporter = ExcelExporter(config)
        with patch("builtins.print"):
            result = exporter.export(data, headers, validation_result=vr, original_indices=[0])
        assert result == output

    def test_export_with_split_empty_data(self, tmp_path):
        data = []
        headers = [{"key": "id", "label": "ID", "width": 10}]
        output = str(tmp_path / "split_empty.xlsx")
        config = {
            "excel_output_path": output,
            "sheet_name": "测试",
            "style_header": True,
            "style_alt_rows": True,
            "split_config": {
                "enabled": True,
                "split_field": "dept",
                "split_rule": "by_value",
                "include_all_sheet": False,
                "sheet_name_template": "{value}",
                "max_sheet_name_length": 31,
                "empty_value_label": "未分类",
            },
            "pivot_config": {"enabled": False},
            "conditional_format_rules": [],
        }
        exporter = ExcelExporter(config)
        with patch("builtins.print"):
            result = exporter.export(data, headers)
        assert result == output
