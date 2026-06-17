import os
import sys
import unittest
import tempfile
import shutil
import json
from unittest import mock
from io import StringIO

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from json_to_excel import (
    _match_custom_rule,
    _match_range_group,
    split_data_by_field,
    _to_numeric,
    _aggregate_values,
    _get_field_label,
    _get_row_key,
    _get_col_key,
    build_pivot_table,
    _compute_pivot_value,
    _compute_row_total,
    _compute_col_total,
    _compute_grand_total,
    _format_value,
    _get_aggregate_label,
    set_column_widths,
    _write_sheet_data,
    apply_header_style,
    apply_data_style,
    apply_conditional_formatting,
    export_to_excel,
    export_to_excel_with_split,
    _apply_validation_marks,
    _print_validation_errors,
)

from openpyxl import Workbook
from config_manager import get_default_config


class TestMatchCustomRule(unittest.TestCase):
    def test_values_list_match(self):
        rule = {"values": ["a", "b", "c"]}
        self.assertTrue(_match_custom_rule("a", rule))
        self.assertFalse(_match_custom_rule("d", rule))

    def test_values_single_match(self):
        rule = {"values": "test"}
        self.assertTrue(_match_custom_rule("test", rule))
        self.assertFalse(_match_custom_rule("other", rule))

    def test_condition_match_true(self):
        rule = {"condition": "value > 10"}
        self.assertTrue(_match_custom_rule(15, rule))
        self.assertFalse(_match_custom_rule(5, rule))

    def test_condition_exception_returns_false(self):
        rule = {"condition": "value + "}
        self.assertFalse(_match_custom_rule(10, rule))

    def test_min_max_both_inclusive(self):
        rule = {"min": 0, "max": 100, "include_min": True, "include_max": True}
        self.assertTrue(_match_custom_rule(0, rule))
        self.assertTrue(_match_custom_rule(50, rule))
        self.assertTrue(_match_custom_rule(100, rule))

    def test_min_max_both_exclusive(self):
        rule = {"min": 0, "max": 100, "include_min": False, "include_max": False}
        self.assertFalse(_match_custom_rule(0, rule))
        self.assertTrue(_match_custom_rule(50, rule))
        self.assertFalse(_match_custom_rule(100, rule))

    def test_min_only(self):
        rule = {"min": 10}
        self.assertTrue(_match_custom_rule(10, rule))
        self.assertTrue(_match_custom_rule(20, rule))
        self.assertFalse(_match_custom_rule(5, rule))

    def test_max_only(self):
        rule = {"max": 100, "include_max": True}
        self.assertTrue(_match_custom_rule(100, rule))
        self.assertTrue(_match_custom_rule(50, rule))
        self.assertFalse(_match_custom_rule(150, rule))

    def test_max_only_exclusive(self):
        rule = {"max": 100}
        self.assertFalse(_match_custom_rule(100, rule))
        self.assertTrue(_match_custom_rule(99, rule))

    def test_none_value_treated_as_zero(self):
        rule = {"min": 0, "max": 10}
        self.assertTrue(_match_custom_rule(None, rule))
        self.assertTrue(_match_custom_rule("", rule))

    def test_invalid_value_returns_false(self):
        rule = {"min": 0, "max": 100}
        self.assertFalse(_match_custom_rule("not_a_number", rule))

    def test_no_condition_returns_false(self):
        rule = {}
        self.assertFalse(_match_custom_rule("any", rule))


class TestMatchRangeGroup(unittest.TestCase):
    def test_both_inclusive(self):
        group = {"min": 0, "max": 100, "include_min": True, "include_max": True}
        self.assertTrue(_match_range_group(0, group))
        self.assertTrue(_match_range_group(50, group))
        self.assertTrue(_match_range_group(100, group))

    def test_both_exclusive(self):
        group = {"min": 0, "max": 100, "include_min": False, "include_max": False}
        self.assertFalse(_match_range_group(0, group))
        self.assertTrue(_match_range_group(50, group))
        self.assertFalse(_match_range_group(100, group))

    def test_min_only(self):
        group = {"min": 10}
        self.assertTrue(_match_range_group(10, group))
        self.assertTrue(_match_range_group(20, group))
        self.assertFalse(_match_range_group(5, group))

    def test_max_only(self):
        group = {"max": 100, "include_max": True}
        self.assertTrue(_match_range_group(100, group))
        self.assertTrue(_match_range_group(50, group))
        self.assertFalse(_match_range_group(150, group))

    def test_max_only_exclusive(self):
        group = {"max": 100}
        self.assertFalse(_match_range_group(100, group))
        self.assertTrue(_match_range_group(99, group))

    def test_none_value_returns_false(self):
        group = {"min": 0, "max": 100}
        self.assertFalse(_match_range_group(None, group))
        self.assertFalse(_match_range_group("", group))

    def test_invalid_value_returns_false(self):
        group = {"min": 0, "max": 100}
        self.assertFalse(_match_range_group("not_a_number", group))


class TestSplitDataByField(unittest.TestCase):
    def setUp(self):
        self.data = [
            {"name": "张三", "dept": "技术部", "salary": 10000},
            {"name": "李四", "dept": "市场部", "salary": 8000},
            {"name": "王五", "dept": "技术部", "salary": 15000},
            {"name": "赵六", "dept": "", "salary": 6000},
        ]

    def test_split_by_value(self):
        config = {"split_rule": "by_value", "empty_value_label": "未分类"}
        groups = split_data_by_field(self.data, "dept", config)
        self.assertEqual(len(groups), 3)
        group_names = [g["name"] for g in groups]
        self.assertIn("技术部", group_names)
        self.assertIn("市场部", group_names)
        self.assertIn("未分类", group_names)

    def test_split_by_range(self):
        config = {
            "split_rule": "by_range",
            "range_groups": [
                {"name": "低薪", "max": 8000, "include_max": True},
                {"name": "中薪", "min": 8000, "max": 12000, "include_min": False, "include_max": True},
                {"name": "高薪", "min": 12000, "include_min": False},
            ]
        }
        groups = split_data_by_field(self.data, "salary", config)
        group_names = [g["name"] for g in groups]
        self.assertIn("低薪", group_names)
        self.assertIn("中薪", group_names)
        self.assertIn("高薪", group_names)

    def test_split_by_custom(self):
        config = {
            "split_rule": "by_custom",
            "custom_rules": [
                {"name": "技术", "condition": "dept == '技术部'"},
                {"name": "非技术", "values": ["市场部", "销售部"]},
            ],
            "other_group_name": "其他",
            "include_other_group": True,
        }
        groups = split_data_by_field(self.data, "dept", config)
        self.assertTrue(len(groups) > 0)

    def test_empty_data(self):
        config = {"split_rule": "by_value"}
        groups = split_data_by_field([], "dept", config)
        self.assertEqual(len(groups), 0)

    def test_original_indices(self):
        config = {"split_rule": "by_value"}
        groups = split_data_by_field(self.data, "dept", config)
        for g in groups:
            self.assertIn("data", g)
            self.assertIn("original_indices", g)
            self.assertEqual(len(g["data"]), len(g["original_indices"]))

    def test_invalid_rule_returns_empty(self):
        config = {"split_rule": "invalid_rule"}
        groups = split_data_by_field(self.data, "dept", config)
        self.assertEqual(len(groups), 0)


class TestToNumeric(unittest.TestCase):
    def test_none_value(self):
        self.assertIsNone(_to_numeric(None))

    def test_empty_string(self):
        self.assertIsNone(_to_numeric(""))

    def test_integer_string(self):
        self.assertEqual(_to_numeric("42"), 42.0)

    def test_float_string(self):
        self.assertEqual(_to_numeric("3.14"), 3.14)

    def test_negative_number(self):
        self.assertEqual(_to_numeric("-100"), -100.0)

    def test_invalid_string(self):
        self.assertIsNone(_to_numeric("abc"))

    def test_integer(self):
        self.assertEqual(_to_numeric(42), 42.0)

    def test_float(self):
        self.assertEqual(_to_numeric(3.14), 3.14)


class TestAggregateValues(unittest.TestCase):
    def setUp(self):
        self.values = [10, 20, 30, 40, 50]
        self.mixed_values = [10, "abc", None, 20, ""]

    def test_sum(self):
        self.assertEqual(_aggregate_values(self.values, "sum"), 150)

    def test_count(self):
        self.assertEqual(_aggregate_values(self.values, "count"), 5)

    def test_count_num(self):
        self.assertEqual(_aggregate_values(self.mixed_values, "count_num"), 2)

    def test_average(self):
        self.assertEqual(_aggregate_values(self.values, "average"), 30.0)

    def test_max(self):
        self.assertEqual(_aggregate_values(self.values, "max"), 50)

    def test_min(self):
        self.assertEqual(_aggregate_values(self.values, "min"), 10)

    def test_product(self):
        self.assertEqual(_aggregate_values([2, 3, 4], "product"), 24)

    def test_stddev(self):
        result = _aggregate_values([2, 4, 6], "stddev")
        self.assertAlmostEqual(result, 2.0, places=2)

    def test_stddevp(self):
        result = _aggregate_values([2, 4, 6], "stddevp")
        self.assertAlmostEqual(result, 1.63, places=2)

    def test_var(self):
        result = _aggregate_values([2, 4, 6], "var")
        self.assertAlmostEqual(result, 4.0, places=2)

    def test_varp(self):
        result = _aggregate_values([2, 4, 6], "varp")
        self.assertAlmostEqual(result, 2.67, places=2)

    def test_default_sum(self):
        self.assertEqual(_aggregate_values(self.values, "unknown"), 150)

    def test_empty_values(self):
        self.assertEqual(_aggregate_values([], "sum"), 0)
        self.assertEqual(_aggregate_values([], "count"), 0)
        self.assertEqual(_aggregate_values([], "average"), 0)
        self.assertEqual(_aggregate_values([], "max"), 0)
        self.assertEqual(_aggregate_values([], "min"), 0)
        self.assertEqual(_aggregate_values([], "product"), 0)
        self.assertEqual(_aggregate_values([], "stddev"), 0)
        self.assertEqual(_aggregate_values([], "stddevp"), 0)
        self.assertEqual(_aggregate_values([], "var"), 0)
        self.assertEqual(_aggregate_values([], "varp"), 0)

    def test_single_value_stddev(self):
        self.assertEqual(_aggregate_values([5], "stddev"), 0)
        self.assertEqual(_aggregate_values([5], "var"), 0)


class TestPivotTableHelpers(unittest.TestCase):
    def setUp(self):
        self.headers = [
            {"key": "dept", "label": "部门"},
            {"key": "name", "label": "姓名"},
            {"key": "salary", "label": "薪资"},
        ]

    def test_get_field_label_found(self):
        label = _get_field_label(self.headers, "dept")
        self.assertEqual(label, "部门")

    def test_get_field_label_not_found(self):
        label = _get_field_label(self.headers, "nonexistent")
        self.assertEqual(label, "nonexistent")

    def test_get_row_key_single_field(self):
        item = {"dept": "技术部"}
        key = _get_row_key(item, ["dept"], "(空白)")
        self.assertEqual(key, ("技术部",))

    def test_get_row_key_multiple_fields(self):
        item = {"dept": "技术部", "name": "张三"}
        key = _get_row_key(item, ["dept", "name"], "(空白)")
        self.assertEqual(key, ("技术部", "张三"))

    def test_get_row_key_empty(self):
        item = {}
        key = _get_row_key(item, ["dept"], "(空白)")
        self.assertEqual(key, ("(空白)",))

    def test_get_col_key_single_field(self):
        item = {"dept": "市场部"}
        key = _get_col_key(item, ["dept"], "(空白)")
        self.assertEqual(key, ("市场部",))

    def test_get_col_key_multiple_fields(self):
        item = {"dept": "市场部", "level": "高级"}
        key = _get_col_key(item, ["dept", "level"], "(空白)")
        self.assertEqual(key, ("市场部", "高级"))

    def test_get_col_key_empty(self):
        item = {}
        key = _get_col_key(item, ["dept"], "(空白)")
        self.assertEqual(key, ("(空白)",))

    def test_get_col_key_no_fields(self):
        item = {"dept": "市场部"}
        key = _get_col_key(item, [], "(空白)")
        self.assertEqual(key, ("",))


class TestBuildPivotTable(unittest.TestCase):
    def setUp(self):
        self.data = [
            {"dept": "技术部", "gender": "男", "salary": 10000},
            {"dept": "技术部", "gender": "女", "salary": 12000},
            {"dept": "市场部", "gender": "男", "salary": 8000},
            {"dept": "市场部", "gender": "女", "salary": 9000},
            "not_a_dict",
        ]
        self.headers = [
            {"key": "dept", "label": "部门"},
            {"key": "gender", "label": "性别"},
            {"key": "salary", "label": "薪资"},
        ]

    def test_basic_pivot(self):
        config = {
            "row_fields": ["dept"],
            "column_fields": ["gender"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        result = build_pivot_table(self.data, config, self.headers)
        self.assertIn("row_keys", result)
        self.assertIn("col_keys", result)
        self.assertIn("data", result)
        self.assertEqual(len(result["row_keys"]), 2)
        self.assertEqual(len(result["col_keys"]), 2)

    def test_pivot_with_empty_values(self):
        config = {
            "row_fields": ["dept"],
            "column_fields": ["gender"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        result = build_pivot_table([], config, self.headers)
        self.assertEqual(len(result["row_keys"]), 0)
        self.assertEqual(len(result["col_keys"]), 0)

    def test_compute_pivot_value(self):
        config = {
            "row_fields": ["dept"],
            "column_fields": ["gender"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        result = build_pivot_table(self.data, config, self.headers)
        vf = {"field": "salary", "aggregate": "sum"}
        value = _compute_pivot_value(result, ("技术部",), ("男",), vf)
        self.assertEqual(value, 10000)

    def test_compute_pivot_value_missing(self):
        config = {
            "row_fields": ["dept"],
            "column_fields": ["gender"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        result = build_pivot_table(self.data, config, self.headers)
        vf = {"field": "salary", "aggregate": "sum"}
        value = _compute_pivot_value(result, ("nonexistent",), ("男",), vf)
        self.assertEqual(value, 0)

    def test_compute_row_total_sum(self):
        config = {
            "row_fields": ["dept"],
            "column_fields": ["gender"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        result = build_pivot_table(self.data, config, self.headers)
        vf = {"field": "salary", "aggregate": "sum"}
        total = _compute_row_total(result, ("技术部",), vf)
        self.assertEqual(total, 22000)

    def test_compute_col_total_sum(self):
        config = {
            "row_fields": ["dept"],
            "column_fields": ["gender"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        result = build_pivot_table(self.data, config, self.headers)
        vf = {"field": "salary", "aggregate": "sum"}
        total = _compute_col_total(result, ("男",), vf)
        self.assertEqual(total, 18000)

    def test_compute_grand_total_sum(self):
        config = {
            "row_fields": ["dept"],
            "column_fields": ["gender"],
            "value_fields": [{"field": "salary", "aggregate": "sum"}],
        }
        result = build_pivot_table(self.data, config, self.headers)
        vf = {"field": "salary", "aggregate": "sum"}
        total = _compute_grand_total(result, vf)
        self.assertEqual(total, 39000)

    def test_format_value_sum(self):
        self.assertEqual(_format_value(100, "sum"), 100)

    def test_format_value_average(self):
        result = _format_value(3.14159, "average")
        self.assertIsInstance(result, float)

    def test_get_aggregate_label_sum(self):
        self.assertEqual(_get_aggregate_label("sum"), "求和")

    def test_get_aggregate_label_count(self):
        self.assertEqual(_get_aggregate_label("count"), "计数")

    def test_get_aggregate_label_average(self):
        self.assertEqual(_get_aggregate_label("average"), "平均值")

    def test_get_aggregate_label_max(self):
        self.assertEqual(_get_aggregate_label("max"), "最大值")

    def test_get_aggregate_label_min(self):
        self.assertEqual(_get_aggregate_label("min"), "最小值")

    def test_get_aggregate_label_unknown(self):
        self.assertEqual(_get_aggregate_label("unknown"), "unknown")


class TestSetColumnWidths(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active

    def test_set_column_widths(self):
        headers = [
            {"key": "name", "label": "姓名", "width": 15},
            {"key": "age", "label": "年龄"},
        ]
        for h in headers:
            self.ws.append([h["label"]])
        set_column_widths(self.ws, headers)
        self.assertEqual(self.ws.column_dimensions["A"].width, 15)

    def test_default_width(self):
        headers = [{"key": "name", "label": "姓名"}]
        self.ws.append(["姓名"])
        set_column_widths(self.ws, headers)
        self.assertIsNotNone(self.ws.column_dimensions["A"].width)


class TestApplyHeaderStyle(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.append(["姓名", "年龄"])
        self.config = get_default_config()

    def test_apply_header_style(self):
        header_cells = self.ws[1]
        apply_header_style(self.ws, header_cells, self.config)
        for cell in header_cells:
            self.assertTrue(cell.font.bold)
            self.assertEqual(cell.alignment.horizontal, "center")

    def test_apply_header_style_no_border(self):
        config = get_default_config()
        config["header_style"]["border_style"] = None
        header_cells = self.ws[1]
        apply_header_style(self.ws, header_cells, config)
        self.assertIsNotNone(header_cells[0].font)


class TestApplyDataStyle(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.headers = [{"key": "name", "label": "姓名"}, {"key": "age", "label": "年龄"}]
        self.ws.append(["姓名", "年龄"])
        self.ws.append(["张三", 30])
        self.ws.append(["李四", 25])
        self.config = get_default_config()

    def test_apply_data_style(self):
        apply_data_style(self.ws, self.headers, 2, self.config)
        cell = self.ws.cell(row=2, column=1)
        self.assertEqual(cell.alignment.horizontal, "left")

    def test_apply_data_style_alt_rows(self):
        config = get_default_config()
        config["style_alt_rows"] = True
        apply_data_style(self.ws, self.headers, 2, config)
        cell_even = self.ws.cell(row=2, column=1)
        cell_odd = self.ws.cell(row=3, column=1)
        self.assertIsNotNone(cell_even.fill)
        self.assertIsNotNone(cell_odd.fill)

    def test_apply_data_style_no_alt_rows(self):
        config = get_default_config()
        config["style_alt_rows"] = False
        apply_data_style(self.ws, self.headers, 2, config)
        cell = self.ws.cell(row=2, column=1)
        self.assertIsNotNone(cell.font)


class TestApplyConditionalFormatting(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.headers = [{"key": "score", "label": "分数"}]
        self.data = [{"score": 95}, {"score": 60}, {"score": 40}]
        self.ws.append(["分数"])
        for d in self.data:
            self.ws.append([d["score"]])
        self.config = get_default_config()

    def test_no_rules_no_change(self):
        apply_conditional_formatting(self.ws, self.headers, self.data, self.config)
        cell = self.ws.cell(row=2, column=1)
        self.assertIsNotNone(cell)

    def test_with_rules(self):
        rules = [
            {
                "field": "score",
                "rule_type": "numeric",
                "operator": "gt",
                "value": 90,
                "style": {
                    "font_bold": True,
                    "bg_color": "#FF0000",
                },
            }
        ]
        self.config["conditional_format_rules"] = rules
        apply_conditional_formatting(self.ws, self.headers, self.data, self.config)
        cell = self.ws.cell(row=2, column=1)
        self.assertTrue(cell.font.bold)


class TestWriteSheetData(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.headers = [
            {"key": "name", "label": "姓名"},
            {"key": "age", "label": "年龄"},
        ]
        self.data = [
            {"name": "张三", "age": 30},
            {"name": "李四", "age": 25},
        ]
        self.config = get_default_config()

    def test_write_basic_data(self):
        _write_sheet_data(self.ws, self.data, self.headers, self.config)
        self.assertEqual(self.ws.cell(row=1, column=1).value, "姓名")
        self.assertEqual(self.ws.cell(row=2, column=1).value, "张三")
        self.assertEqual(self.ws.cell(row=3, column=1).value, "李四")

    def test_write_with_progress(self):
        from json_to_excel import ProgressTracker
        progress = ProgressTracker(total=2, description="test", unit="行")
        _write_sheet_data(self.ws, self.data, self.headers, self.config, progress=progress)
        self.assertEqual(self.ws.max_row, 3)

    def test_write_with_computed_cache(self):
        computed_cache = {}
        item = self.data[0]
        computed_cache[id(item)] = {"name": "张三(计算)", "age": 31}
        _write_sheet_data(self.ws, self.data, self.headers, self.config, computed_cache=computed_cache)
        self.assertEqual(self.ws.cell(row=2, column=1).value, "张三(计算)")

    def test_write_with_invalid_items(self):
        data = ["not_a_dict", {"name": "张三", "age": 30}]
        _write_sheet_data(self.ws, data, self.headers, self.config)
        self.assertEqual(self.ws.cell(row=2, column=1).value, "张三")
        self.assertIsNone(self.ws.cell(row=3, column=1).value)

    def test_freeze_panes(self):
        _write_sheet_data(self.ws, self.data, self.headers, self.config)
        self.assertEqual(self.ws.freeze_panes, "A2")


class TestExportToExcel(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.headers = [
            {"key": "name", "label": "姓名"},
            {"key": "age", "label": "年龄"},
        ]
        self.data = [
            {"name": "张三", "age": 30},
            {"name": "李四", "age": 25},
        ]
        self.config = get_default_config()
        self.output_path = os.path.join(self.temp_dir, "test_output.xlsx")
        self.config["excel_output_path"] = self.output_path

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_basic(self):
        result = export_to_excel(self.data, self.headers, self.config)
        self.assertEqual(result, self.output_path)
        self.assertTrue(os.path.exists(self.output_path))

    def test_export_with_sheet_name(self):
        self.config["sheet_name"] = "员工数据"
        result = export_to_excel(self.data, self.headers, self.config)
        self.assertTrue(os.path.exists(result))

    def test_export_empty_data(self):
        result = export_to_excel([], self.headers, self.config)
        self.assertTrue(os.path.exists(result))

    def test_export_creates_directory(self):
        nested_path = os.path.join(self.temp_dir, "subdir", "test.xlsx")
        self.config["excel_output_path"] = nested_path
        result = export_to_excel(self.data, self.headers, self.config)
        self.assertTrue(os.path.exists(result))

    def test_export_with_pivot_table(self):
        self.config["pivot_tables"] = [
            {
                "sheet_name": "数据透视表",
                "row_fields": ["name"],
                "value_fields": [{"field": "age", "aggregate": "sum"}],
            }
        ]
        result = export_to_excel(self.data, self.headers, self.config)
        self.assertTrue(os.path.exists(result))


class TestExportToExcelWithSplit(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.headers = [
            {"key": "name", "label": "姓名"},
            {"key": "dept", "label": "部门"},
            {"key": "age", "label": "年龄"},
        ]
        self.data = [
            {"name": "张三", "dept": "技术部", "age": 30},
            {"name": "李四", "dept": "市场部", "age": 25},
            {"name": "王五", "dept": "技术部", "age": 35},
        ]
        self.config = get_default_config()
        self.output_path = os.path.join(self.temp_dir, "test_split.xlsx")
        self.config["excel_output_path"] = self.output_path
        self.config["split_config"] = {
            "split_field": "dept",
            "split_rule": "by_value",
            "empty_value_label": "未分类",
            "include_all_sheet": True,
            "all_sheet_name": "全部数据",
        }

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_with_split(self):
        result = export_to_excel_with_split(self.data, self.headers, self.config)
        self.assertEqual(result, self.output_path)
        self.assertTrue(os.path.exists(result))


class TestValidationMarks(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.headers = [{"key": "name", "label": "姓名"}, {"key": "email", "label": "邮箱"}]
        self.ws.append(["姓名", "邮箱"])
        self.ws.append(["张三", "invalid"])
        self.original_indices = [0, 1]

    def test_apply_validation_marks(self):
        from data_validator import ValidationResult, ValidationError, ValidationRule, VALIDATION_TYPE_FORMAT, FORMAT_EMAIL, ON_FAIL_MARK
        result = ValidationResult()
        rule = ValidationRule("email", VALIDATION_TYPE_FORMAT, {"format": FORMAT_EMAIL}, on_fail=ON_FAIL_MARK)
        error = ValidationError(0, rule, "invalid", "email")
        result.add_error(error)
        _apply_validation_marks(self.ws, result, self.headers, self.original_indices)
        self.assertIsNotNone(self.ws.cell(row=2, column=2).fill)


class TestPrintValidationErrors(unittest.TestCase):
    def test_print_errors(self):
        from data_validator import ValidationResult, ValidationError, ValidationRule, VALIDATION_TYPE_NOT_NULL, ON_FAIL_MARK
        result = ValidationResult()
        rule = ValidationRule("name", VALIDATION_TYPE_NOT_NULL, on_fail=ON_FAIL_MARK)
        error = ValidationError(0, rule, None, "name")
        result.add_error(error)
        data = [{"name": None}]
        with mock.patch("builtins.print"):
            _print_validation_errors(result, data)

    def test_print_no_errors(self):
        from data_validator import ValidationResult
        result = ValidationResult()
        data = []
        with mock.patch("builtins.print") as mock_print:
            _print_validation_errors(result, data)
            mock_print.assert_not_called()
