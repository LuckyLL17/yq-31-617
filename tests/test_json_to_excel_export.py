import unittest
import sys
import os
import tempfile
import shutil
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from json_to_excel import (
    export_to_excel,
    export_to_excel_with_split,
    _apply_validation_marks,
    _print_validation_errors,
    _sanitize_sheet_name,
    _render_sheet_name,
    ProgressTracker,
)
from data_validator import ValidationRule, ValidationResult, ValidationError


SAMPLE_DATA = [
    {"name": "Alice", "age": 30, "email": "alice@example.com", "department": "Engineering"},
    {"name": "Bob", "age": 25, "email": "bob@example.com", "department": "Marketing"},
    {"name": "Charlie", "age": 35, "email": "charlie@example.com", "department": "Engineering"},
]

SAMPLE_HEADERS = [
    {"key": "name", "label": "姓名"},
    {"key": "age", "label": "年龄"},
    {"key": "email", "label": "邮箱"},
    {"key": "department", "label": "部门"},
]


class TestExportToExcelBasic(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.temp_dir, "test.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_basic(self):
        config = {"excel_output_path": self.output_path}
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertEqual(result, self.output_path)
        self.assertTrue(os.path.exists(result))

    def test_export_with_sheet_name(self):
        config = {
            "excel_output_path": self.output_path,
            "sheet_name": "用户数据",
        }
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)

    def test_export_new_directory(self):
        new_dir = os.path.join(self.temp_dir, "subdir", "nested")
        output = os.path.join(new_dir, "test.xlsx")
        config = {"excel_output_path": output}
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)
        self.assertTrue(os.path.exists(result))

    def test_export_no_style_header(self):
        config = {
            "excel_output_path": self.output_path,
            "style_header": False,
        }
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)


class TestExportToExcelWithValidation(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.temp_dir, "test.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_with_validation_mark(self):
        rules = [
            {
                "field": "email",
                "rule_type": "format",
                "params": {"format": "email"},
                "message": "邮箱格式不正确",
                "on_fail": "mark",
            }
        ]
        data = [
            {"name": "Alice", "age": 30, "email": "invalid-email"},
            {"name": "Bob", "age": 25, "email": "bob@example.com"},
        ]
        config = {
            "excel_output_path": self.output_path,
            "validation_rules": rules,
        }
        result = export_to_excel(data, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)

    def test_export_with_validation_skip(self):
        rules = [
            {
                "field": "email",
                "rule_type": "format",
                "params": {"format": "email"},
                "message": "邮箱格式不正确",
                "on_fail": "skip",
            }
        ]
        data = [
            {"name": "Alice", "age": 30, "email": "invalid-email"},
            {"name": "Bob", "age": 25, "email": "bob@example.com"},
        ]
        config = {
            "excel_output_path": self.output_path,
            "validation_rules": rules,
        }
        result = export_to_excel(data, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)

    def test_export_validation_aborted(self):
        rules = [
            {
                "field": "age",
                "rule_type": "not_null",
                "params": {},
                "message": "年龄不能为空",
                "on_fail": "abort",
            }
        ]
        data = [
            {"name": "Alice", "age": None, "email": "alice@example.com"},
        ]
        config = {
            "excel_output_path": self.output_path,
            "validation_rules": rules,
        }
        result = export_to_excel(data, SAMPLE_HEADERS, config)
        self.assertIsNone(result)


class TestExportToExcelWithComputedColumns(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.temp_dir, "test.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_with_computed_columns(self):
        computed_columns = [
            {
                "key": "full_info",
                "label": "完整信息",
                "type": "template",
                "template": "{name} - {age}岁",
                "enabled": True,
            }
        ]
        config = {
            "excel_output_path": self.output_path,
            "computed_columns": computed_columns,
        }
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)

    def test_export_with_disabled_computed_columns(self):
        computed_columns = [
            {
                "key": "full_info",
                "label": "完整信息",
                "type": "template",
                "template": "{name} - {age}岁",
                "enabled": False,
            }
        ]
        config = {
            "excel_output_path": self.output_path,
            "computed_columns": computed_columns,
        }
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)


class TestExportToExcelWithSplit(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.temp_dir, "test.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_with_split_by_value(self):
        split_config = {
            "enabled": True,
            "split_field": "department",
            "split_rule": "by_value",
            "include_all_sheet": True,
        }
        config = {
            "excel_output_path": self.output_path,
            "split_config": split_config,
        }
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)

    def test_export_with_split_no_all_sheet(self):
        split_config = {
            "enabled": True,
            "split_field": "department",
            "split_rule": "by_value",
            "include_all_sheet": False,
        }
        config = {
            "excel_output_path": self.output_path,
            "split_config": split_config,
        }
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)


class TestExportToExcelWithSplitAdvanced(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.temp_dir, "test.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_with_split_custom_template(self):
        split_config = {
            "enabled": True,
            "split_field": "department",
            "split_rule": "by_value",
            "sheet_name_template": "部门-{value}",
            "include_all_sheet": True,
            "all_sheet_name": "全部",
        }
        config = {
            "excel_output_path": self.output_path,
            "split_config": split_config,
        }
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)

    def test_export_with_split_duplicate_names(self):
        data = [
            {"name": "A", "group": "Team"},
            {"name": "B", "group": "Team"},
            {"name": "C", "group": "Team_2"},
        ]
        headers = [
            {"key": "name", "label": "姓名"},
            {"key": "group", "label": "分组"},
        ]
        split_config = {
            "enabled": True,
            "split_field": "group",
            "split_rule": "by_value",
            "sheet_name_template": "Team",
            "max_sheet_name_length": 10,
            "include_all_sheet": False,
        }
        config = {
            "excel_output_path": self.output_path,
            "split_config": split_config,
        }
        result = export_to_excel(data, headers, config)
        self.assertIsNotNone(result)


class TestExportToExcelWithNonDictItems(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.temp_dir, "test.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_with_non_dict_items(self):
        data = ["not_dict", {"name": "Alice", "age": 30}, None, {"name": "Bob", "age": 25}]
        headers = [
            {"key": "name", "label": "姓名"},
            {"key": "age", "label": "年龄"},
        ]
        config = {"excel_output_path": self.output_path}
        result = export_to_excel(data, headers, config)
        self.assertIsNotNone(result)


class TestExportExcelWithSplitDirect(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.temp_dir, "test_split.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_split_export_direct_call(self):
        split_config = {
            "enabled": True,
            "split_field": "department",
            "split_rule": "by_value",
        }
        config = {
            "excel_output_path": self.output_path,
            "split_config": split_config,
        }
        result = export_to_excel_with_split(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)

    def test_split_export_with_validation(self):
        rules = [
            {
                "field": "email",
                "rule_type": "format",
                "params": {"format": "email"},
                "message": "邮箱格式错误",
                "on_fail": "mark",
            }
        ]
        split_config = {
            "enabled": True,
            "split_field": "department",
            "split_rule": "by_value",
        }
        config = {
            "excel_output_path": self.output_path,
            "split_config": split_config,
            "validation_rules": rules,
        }
        result = export_to_excel_with_split(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)

    def test_split_export_empty_groups(self):
        data = []
        split_config = {
            "enabled": True,
            "split_field": "department",
            "split_rule": "by_value",
            "include_all_sheet": False,
        }
        config = {
            "excel_output_path": self.output_path,
            "split_config": split_config,
        }
        result = export_to_excel_with_split(data, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)


class TestApplyValidationMarks(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook
        self.wb = Workbook()
        self.ws = self.wb.active

    def test_apply_marks_basic(self):
        validation_result = ValidationResult()
        rule = ValidationRule(
            field="email",
            rule_type="format",
            params={"format": "email"},
            message="邮箱格式不正确",
            on_fail="mark",
        )
        error = ValidationError(row_index=0, field_key="email", value="bad", rule=rule)
        validation_result.add_error(error)
        validation_result.marked_rows.add(0)

        headers = [{"key": "name", "label": "姓名"}, {"key": "email", "label": "邮箱"}]
        original_indices = [0, 1]

        self.ws.append(["姓名", "邮箱"])
        self.ws.append(["Alice", "bad"])
        self.ws.append(["Bob", "bob@example.com"])

        _apply_validation_marks(self.ws, validation_result, headers, original_indices)

        cell = self.ws.cell(row=2, column=1)
        self.assertIsNotNone(cell.font)
        self.assertIsNotNone(cell.fill)

    def test_apply_marks_not_in_original_indices(self):
        validation_result = ValidationResult()
        rule = ValidationRule(
            field="email",
            rule_type="format",
            params={"format": "email"},
            message="邮箱格式错误",
            on_fail="mark",
        )
        error = ValidationError(row_index=5, field_key="email", value="x", rule=rule)
        validation_result.add_error(error)
        validation_result.marked_rows.add(5)

        headers = [{"key": "name", "label": "姓名"}]
        original_indices = [0, 1]

        self.ws.append(["姓名"])
        self.ws.append(["Alice"])
        self.ws.append(["Bob"])

        _apply_validation_marks(self.ws, validation_result, headers, original_indices)

    def test_apply_marks_long_comment(self):
        validation_result = ValidationResult()
        long_msg = "错误信息" * 50
        rule = ValidationRule(
            field="email",
            rule_type="format",
            params={"format": "email"},
            message=long_msg,
            on_fail="mark",
        )
        error = ValidationError(row_index=0, field_key="email", value="bad", rule=rule)
        validation_result.add_error(error)
        validation_result.marked_rows.add(0)

        headers = [{"key": "email", "label": "邮箱"}]
        original_indices = [0]

        self.ws.append(["邮箱"])
        self.ws.append(["bad"])

        _apply_validation_marks(self.ws, validation_result, headers, original_indices)

        cell = self.ws.cell(row=2, column=1)
        self.assertIsNotNone(cell.comment)
        self.assertTrue(len(str(cell.comment.text)) <= 205)


class TestPrintValidationErrors(unittest.TestCase):
    def test_print_errors_basic(self):
        validation_result = ValidationResult()
        rule = ValidationRule(
            field="email",
            rule_type="format",
            params={"format": "email"},
            message="邮箱格式错误",
            on_fail="mark",
        )
        error = ValidationError(row_index=0, field_key="email", value="bad@", rule=rule)
        validation_result.add_error(error)

        data = [{"email": "bad@"}]

        with mock.patch("builtins.print") as mock_print:
            _print_validation_errors(validation_result, data, max_display=10)
            mock_print.assert_called()

    def test_print_errors_empty(self):
        validation_result = ValidationResult()
        data = []

        with mock.patch("builtins.print") as mock_print:
            _print_validation_errors(validation_result, data)
            mock_print.assert_not_called()

    def test_print_errors_more_than_max(self):
        validation_result = ValidationResult()
        rule = ValidationRule(
            field="age",
            rule_type="not_null",
            params={},
            message="不能为空",
            on_fail="skip",
        )
        for i in range(15):
            error = ValidationError(row_index=i, field_key="age", value=None, rule=rule)
            validation_result.add_error(error)

        data = [{"age": None}] * 15

        with mock.patch("builtins.print") as mock_print:
            _print_validation_errors(validation_result, data, max_display=10)
            calls = " ".join(str(c) for c in mock_print.call_args_list)
            self.assertIn("还有", calls)

    def test_print_errors_none_value(self):
        validation_result = ValidationResult()
        rule = ValidationRule(
            field="name",
            rule_type="not_null",
            params={},
            message="姓名不能为空",
            on_fail="mark",
        )
        error = ValidationError(row_index=0, field_key="name", value=None, rule=rule)
        validation_result.add_error(error)

        data = [{"name": None}]

        with mock.patch("builtins.print") as mock_print:
            _print_validation_errors(validation_result, data)
            mock_print.assert_called()


class TestExportToExcelWithPivotTable(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.temp_dir, "test.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_with_pivot_table(self):
        pivot_config = {
            "enabled": True,
            "rows": ["department"],
            "values": [{"field": "age", "aggregation": "average", "label": "平均年龄"}],
        }
        config = {
            "excel_output_path": self.output_path,
            "pivot_config": pivot_config,
        }
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)


class TestExportToExcelWithConditionalFormatting(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.temp_dir, "test.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_with_conditional_formatting(self):
        conditional_formats = [
            {
                "field": "age",
                "type": "numeric",
                "operator": "gt",
                "value": 30,
                "style": {"bg_color": "FFCCCC"},
            }
        ]
        config = {
            "excel_output_path": self.output_path,
            "conditional_formats": conditional_formats,
        }
        result = export_to_excel(SAMPLE_DATA, SAMPLE_HEADERS, config)
        self.assertIsNotNone(result)


if __name__ == "__main__":
    unittest.main()
