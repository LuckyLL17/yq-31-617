import os
import sys
import json
import tempfile
import shutil
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from json_to_excel import (
    DataLoader,
    ExcelExporter,
    load_json,
    auto_detect_headers,
    merge_headers,
    extract_value,
    flatten_dict,
    _sanitize_sheet_name,
    _render_sheet_name,
    _get_column_letter,
    _get_alignment,
    _create_border,
)
from config_manager import ExportConfig


SAMPLE_DATA = [
    {"id": 1, "name": "张三", "age": 28, "email": "zhangsan@example.com"},
    {"id": 2, "name": "李四", "age": 35, "email": "lisi@example.com"},
    {"id": 3, "name": "王五", "age": 25, "email": "wangwu@example.com"},
]

SAMPLE_HEADERS = [
    {"key": "id", "label": "ID"},
    {"key": "name", "label": "姓名"},
    {"key": "age", "label": "年龄"},
    {"key": "email", "label": "邮箱"},
]


class TestDataLoaderInit(unittest.TestCase):
    def test_init_with_none(self):
        loader = DataLoader()
        self.assertIsNotNone(loader.config)
        self.assertIsNone(loader._raw_data)

    def test_init_with_dict(self):
        loader = DataLoader(config={"sheet_name": "测试"})
        self.assertIsNotNone(loader.config)
        self.assertEqual(loader.config.sheet_name, "测试")

    def test_init_with_config_object(self):
        config = ExportConfig.from_default()
        loader = DataLoader(config=config)
        self.assertIsNotNone(loader.config)

    def test_from_file_classmethod(self):
        temp_dir = tempfile.mkdtemp()
        try:
            json_path = os.path.join(temp_dir, "test.json")
            with open(json_path, "w", encoding="utf-8") as f:
                json.dump(SAMPLE_DATA, f, ensure_ascii=False)

            loader = DataLoader.from_file(json_path)
            self.assertEqual(len(loader.raw_data), 3)
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_repr(self):
        loader = DataLoader()
        repr_str = repr(loader)
        self.assertIsInstance(repr_str, str)


class TestDataLoaderLoad(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.json_path = os.path.join(self.temp_dir, "test.json")
        with open(self.json_path, "w", encoding="utf-8") as f:
            json.dump(SAMPLE_DATA, f, ensure_ascii=False)

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_load_from_path(self):
        loader = DataLoader()
        data = loader.load(self.json_path)
        self.assertEqual(len(data), 3)
        self.assertEqual(loader._raw_data[0]["name"], "张三")

    def test_load_from_config_path(self):
        config = {"json_file_path": self.json_path}
        loader = DataLoader(config=config)
        data = loader.load()
        self.assertEqual(len(data), 3)

    def test_load_no_path_raises(self):
        loader = DataLoader(config={"json_file_path": ""})
        with self.assertRaises(ValueError):
            loader.load()

    def test_set_data(self):
        loader = DataLoader()
        result = loader.set_data(SAMPLE_DATA)
        self.assertIs(result, loader)
        self.assertEqual(len(loader.raw_data), 3)

    def test_raw_data_raises_if_not_loaded(self):
        loader = DataLoader()
        with self.assertRaises(ValueError):
            _ = loader.raw_data


class TestDataLoaderHeaders(unittest.TestCase):
    def setUp(self):
        self.loader = DataLoader()
        self.loader.set_data(SAMPLE_DATA)

    def test_detect_headers(self):
        headers = self.loader.detect_headers()
        self.assertEqual(len(headers), 4)
        keys = [h["key"] for h in headers]
        self.assertIn("id", keys)
        self.assertIn("name", keys)

    def test_detect_headers_raises_if_no_data(self):
        loader = DataLoader()
        with self.assertRaises(ValueError):
            loader.detect_headers()

    def test_auto_headers_property(self):
        headers = self.loader.auto_headers
        self.assertEqual(len(headers), 4)

    def test_merge_headers(self):
        self.loader.config.default_headers = [
            {"key": "id", "label": "编号"},
            {"key": "name", "label": "名字"},
        ]
        headers = self.loader.merge_headers()
        self.assertGreaterEqual(len(headers), 2)

    def test_headers_property(self):
        headers = self.loader.headers
        self.assertIsInstance(headers, list)

    def test_headers_triggers_merge(self):
        self.loader._auto_headers = None
        self.loader._merged_headers = None
        headers = self.loader.headers
        self.assertIsNotNone(headers)


class TestDataLoaderComputedColumns(unittest.TestCase):
    def setUp(self):
        self.loader = DataLoader()
        self.loader.set_data(SAMPLE_DATA)

    def test_no_computed_columns(self):
        cache, headers = self.loader.apply_computed_columns()
        self.assertIsNone(cache)

    def test_disabled_computed_columns(self):
        self.loader.config.computed_columns = [
            {"key": "double", "label": "双倍", "enabled": False, "formula": "age * 2", "type": "number"},
        ]
        cache, headers = self.loader.apply_computed_columns()
        self.assertIsNone(cache)

    def test_with_computed_columns(self):
        self.loader.config.computed_columns = [
            {"key": "double_age", "label": "双倍年龄", "enabled": True, "formula": "age * 2", "type": "number"},
        ]
        cache, headers = self.loader.apply_computed_columns()
        self.assertIsNotNone(cache)
        self.assertGreater(len(headers), len(SAMPLE_HEADERS))

    def test_computed_cache_property(self):
        self.assertIsNone(self.loader.computed_cache)


class TestDataLoaderValidation(unittest.TestCase):
    def setUp(self):
        self.loader = DataLoader()
        self.loader.set_data(SAMPLE_DATA)

    def test_no_validation_rules(self):
        result = self.loader.validate_data()
        self.assertIsNone(result)

    def test_with_validation_rules(self):
        self.loader.config.validation_rules = [
            {
                "field": "age",
                "rule_type": "range",
                "min": 0,
                "max": 100,
                "enabled": True,
                "error_level": "error",
            },
        ]
        result = self.loader.validate_data()
        self.assertIsNotNone(result)

    def test_validation_result_property(self):
        self.assertIsNone(self.loader.validation_result)

    def test_get_valid_data_no_validation(self):
        valid_data = self.loader.get_valid_data()
        self.assertEqual(len(valid_data), 3)

    def test_original_indices_property(self):
        indices = self.loader.original_indices
        self.assertEqual(len(indices), 3)
        self.assertEqual(indices[0], 0)


class TestDataLoaderPrepareForExport(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.json_path = os.path.join(self.temp_dir, "test.json")
        with open(self.json_path, "w", encoding="utf-8") as f:
            json.dump(SAMPLE_DATA, f, ensure_ascii=False)

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_prepare_for_export_with_data_set(self):
        loader = DataLoader()
        loader.set_data(SAMPLE_DATA)
        result = loader.prepare_for_export()
        self.assertIsNotNone(result)
        self.assertIn("data", result)
        self.assertIn("headers", result)
        self.assertIn("computed_cache", result)
        self.assertIn("validation_result", result)
        self.assertIn("original_indices", result)

    def test_prepare_for_export_loads_file(self):
        loader = DataLoader(config={"json_file_path": self.json_path})
        result = loader.prepare_for_export()
        self.assertIsNotNone(result)
        self.assertEqual(len(result["data"]), 3)


class TestDataLoaderHelperMethods(unittest.TestCase):
    def setUp(self):
        self.loader = DataLoader()

    def test_extract_value_method(self):
        item = {"name": "test"}
        val = self.loader.extract_value(item, "name")
        self.assertEqual(val, "test")

    def test_flatten_dict_method(self):
        d = {"a": {"b": 1}}
        result = self.loader.flatten_dict(d)
        self.assertIn("a.b", result)


class TestExcelExporterInit(unittest.TestCase):
    def test_init_with_none(self):
        exporter = ExcelExporter()
        self.assertIsNotNone(exporter.config)

    def test_init_with_dict(self):
        exporter = ExcelExporter(config={"sheet_name": "测试表"})
        self.assertEqual(exporter.config.sheet_name, "测试表")

    def test_init_with_config_object(self):
        config = ExportConfig.from_default()
        exporter = ExcelExporter(config=config)
        self.assertIsNotNone(exporter.config)

    def test_from_config_classmethod(self):
        config = ExportConfig.from_default()
        exporter = ExcelExporter.from_config(config)
        self.assertIsInstance(exporter, ExcelExporter)


class TestExcelExporterExport(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.exporter = ExcelExporter()
        self.output_path = os.path.join(self.temp_dir, "test.xlsx")
        self.exporter.config.set_output_path(self.output_path, "excel")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_basic(self):
        result = self.exporter.export(SAMPLE_DATA, SAMPLE_HEADERS)
        self.assertTrue(os.path.exists(result))
        self.assertEqual(result, self.output_path)

    def test_export_creates_dir(self):
        nested_path = os.path.join(self.temp_dir, "subdir", "nested", "test.xlsx")
        self.exporter.config.set_output_path(nested_path, "excel")
        result = self.exporter.export(SAMPLE_DATA, SAMPLE_HEADERS)
        self.assertTrue(os.path.exists(result))

    def test_export_with_computed_cache(self):
        computed_cache = {
            id(SAMPLE_DATA[0]): {"extra": "value1"},
            id(SAMPLE_DATA[1]): {"extra": "value2"},
        }
        headers = SAMPLE_HEADERS + [{"key": "extra", "label": "额外"}]
        result = self.exporter.export(SAMPLE_DATA, headers, computed_cache=computed_cache)
        self.assertTrue(os.path.exists(result))

    def test_export_with_split(self):
        self.exporter.config.split_config = {
            "enabled": True,
            "split_field": "age",
            "split_type": "by_value",
        }
        result = self.exporter.export(SAMPLE_DATA, SAMPLE_HEADERS)
        self.assertTrue(os.path.exists(result))

    def test_export_non_dict_items(self):
        data = [SAMPLE_DATA[0], "not a dict", SAMPLE_DATA[1]]
        result = self.exporter.export(data, SAMPLE_HEADERS)
        self.assertTrue(os.path.exists(result))


class TestExcelExporterFromLoader(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.exporter = ExcelExporter()
        self.output_path = os.path.join(self.temp_dir, "from_loader.xlsx")
        self.exporter.config.set_output_path(self.output_path, "excel")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_from_loader(self):
        loader = DataLoader()
        loader.set_data(SAMPLE_DATA)
        result = self.exporter.export_from_loader(loader)
        self.assertTrue(os.path.exists(result))


class TestHelperFunctions(unittest.TestCase):
    def test_sanitize_sheet_name(self):
        name = _sanitize_sheet_name("正常名称")
        self.assertEqual(name, "正常名称")

    def test_sanitize_sheet_name_long(self):
        long_name = "A" * 50
        result = _sanitize_sheet_name(long_name)
        self.assertLessEqual(len(result), 31)

    def test_sanitize_sheet_name_invalid_chars(self):
        name = _sanitize_sheet_name("name[]:*?/\\")
        self.assertNotIn("[", name)
        self.assertNotIn("]", name)

    def test_render_sheet_name(self):
        result = _render_sheet_name("Sheet_{value}", "测试", 1, 5)
        self.assertIn("测试", result)

    def test_render_sheet_name_empty_value(self):
        result = _render_sheet_name("Sheet_{value}", "", 1, 5)
        self.assertIn("未分类", result)

    def test_get_column_letter(self):
        self.assertEqual(_get_column_letter(1), "A")
        self.assertEqual(_get_column_letter(26), "Z")
        self.assertEqual(_get_column_letter(27), "AA")

    def test_get_alignment_center(self):
        align = _get_alignment("center")
        self.assertEqual(align, "center")

    def test_get_alignment_left(self):
        align = _get_alignment("left")
        self.assertEqual(align, "left")

    def test_get_alignment_right(self):
        align = _get_alignment("right")
        self.assertEqual(align, "right")

    def test_get_alignment_default(self):
        align = _get_alignment("unknown")
        self.assertEqual(align, "center")

    def test_create_border(self):
        from openpyxl.styles import Border
        border = _create_border({"border_style": "thin", "border_color": "000000"})
        self.assertIsInstance(border, Border)

    def test_create_border_empty(self):
        from openpyxl.styles import Border
        border = _create_border({})
        self.assertIsInstance(border, Border)


class TestMergeHeaders(unittest.TestCase):
    def test_merge_with_config_headers_auto_detect_true(self):
        config_headers = [
            {"key": "id", "label": "编号"},
            {"key": "name", "label": "姓名"},
        ]
        auto_headers = [
            {"key": "id", "label": "id"},
            {"key": "age", "label": "age"},
        ]
        result = merge_headers(config_headers, auto_headers, {"auto_detect_headers": True})
        self.assertEqual(len(result), 3)
        self.assertEqual(result[0]["label"], "编号")

    def test_merge_with_config_headers_auto_detect_false(self):
        config_headers = [
            {"key": "id", "label": "编号"},
            {"key": "name", "label": "姓名"},
        ]
        auto_headers = [
            {"key": "id", "label": "id"},
            {"key": "age", "label": "age"},
        ]
        result = merge_headers(config_headers, auto_headers, {"auto_detect_headers": False})
        self.assertEqual(len(result), 2)

    def test_merge_empty_config_headers(self):
        auto_headers = [{"key": "id", "label": "ID"}]
        result = merge_headers([], auto_headers, {})
        self.assertEqual(len(result), 1)


if __name__ == "__main__":
    unittest.main()
