#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import unittest
import os
import sys
import tempfile
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from json_to_excel import (
    build_pivot_table,
    _compute_pivot_value,
    _compute_row_total,
    _compute_col_total,
    _compute_grand_total,
    _format_value,
    _get_field_label,
    _get_row_key,
    _get_col_key,
    _aggregate_values,
    _to_numeric,
    _apply_validation_marks,
    _print_validation_errors,
    _match_custom_rule,
    _match_range_group,
    _sanitize_sheet_name,
    _render_sheet_name,
    create_pivot_sheet,
    add_pivot_table_to_workbook,
    get_default_config,
    split_data_by_field,
)

from data_validator import (
    ValidationResult,
    ValidationRule,
    ValidationError,
    ON_FAIL_MARK,
)

from openpyxl import Workbook


SAMPLE_DATA = [
    {"region": "华东", "product": "A", "sales": 100, "quantity": 10},
    {"region": "华东", "product": "B", "sales": 200, "quantity": 20},
    {"region": "华北", "product": "A", "sales": 150, "quantity": 15},
    {"region": "华北", "product": "B", "sales": 250, "quantity": 25},
    {"region": "华南", "product": "A", "sales": 300, "quantity": 30},
]


class TestToNumeric(unittest.TestCase):
    def test_int_string(self):
        self.assertEqual(_to_numeric("123"), 123)

    def test_float_string(self):
        self.assertEqual(_to_numeric("123.45"), 123.45)

    def test_int_value(self):
        self.assertEqual(_to_numeric(123), 123)

    def test_float_value(self):
        self.assertEqual(_to_numeric(123.45), 123.45)

    def test_invalid_string(self):
        self.assertIsNone(_to_numeric("abc"))

    def test_none_value(self):
        self.assertIsNone(_to_numeric(None))


class TestAggregateValues(unittest.TestCase):
    def test_sum(self):
        self.assertEqual(_aggregate_values([1, 2, 3], "sum"), 6)

    def test_average(self):
        self.assertEqual(_aggregate_values([1, 2, 3], "average"), 2)

    def test_count(self):
        self.assertEqual(_aggregate_values([1, None, 3], "count"), 3)

    def test_count_num(self):
        self.assertEqual(_aggregate_values([1, None, "abc", 3], "count_num"), 2)

    def test_max(self):
        self.assertEqual(_aggregate_values([1, 3, 2], "max"), 3)

    def test_min(self):
        self.assertEqual(_aggregate_values([1, 3, 2], "min"), 1)

    def test_product(self):
        self.assertEqual(_aggregate_values([2, 3, 4], "product"), 24)

    def test_stddev(self):
        result = _aggregate_values([1, 2, 3], "stddev")
        self.assertAlmostEqual(result, 1.0, places=1)

    def test_stddevp(self):
        result = _aggregate_values([1, 2, 3], "stddevp")
        self.assertAlmostEqual(result, 0.816, places=2)

    def test_var(self):
        result = _aggregate_values([1, 2, 3], "var")
        self.assertAlmostEqual(result, 1.0, places=1)

    def test_varp(self):
        result = _aggregate_values([1, 2, 3], "varp")
        self.assertAlmostEqual(result, 0.667, places=2)

    def test_empty_sum(self):
        self.assertEqual(_aggregate_values([], "sum"), 0)

    def test_empty_count(self):
        self.assertEqual(_aggregate_values([], "count"), 0)

    def test_empty_average(self):
        self.assertEqual(_aggregate_values([], "average"), 0)

    def test_empty_max(self):
        self.assertEqual(_aggregate_values([], "max"), 0)

    def test_empty_product(self):
        self.assertEqual(_aggregate_values([], "product"), 0)

    def test_stddev_less_than_two(self):
        self.assertEqual(_aggregate_values([1], "stddev"), 0)

    def test_var_less_than_two(self):
        self.assertEqual(_aggregate_values([1], "var"), 0)


class TestGetFieldLabel(unittest.TestCase):
    def test_found(self):
        headers = [{"key": "id", "label": "编号"}, {"key": "name", "label": "姓名"}]
        self.assertEqual(_get_field_label(headers, "name"), "姓名")

    def test_not_found(self):
        headers = [{"key": "id", "label": "编号"}]
        self.assertEqual(_get_field_label(headers, "unknown"), "unknown")


class TestGetRowKey(unittest.TestCase):
    def test_single_field(self):
        item = {"region": "华东", "product": "A"}
        rk = _get_row_key(item, ["region"], "(空白)")
        self.assertEqual(rk, ("华东",))

    def test_multiple_fields(self):
        item = {"region": "华东", "product": "A"}
        rk = _get_row_key(item, ["region", "product"], "(空白)")
        self.assertEqual(rk, ("华东", "A"))

    def test_empty_field_value(self):
        item = {"region": None, "product": "A"}
        rk = _get_row_key(item, ["region", "product"], "(空白)")
        self.assertEqual(rk, ("(空白)", "A"))

    def test_empty_string_field(self):
        item = {"region": "", "product": "A"}
        rk = _get_row_key(item, ["region", "product"], "(空白)")
        self.assertEqual(rk, ("(空白)", "A"))

    def test_no_row_fields(self):
        item = {"region": "华东"}
        rk = _get_row_key(item, [], "(空白)")
        self.assertEqual(rk, ())


class TestGetColKey(unittest.TestCase):
    def test_single_field(self):
        item = {"region": "华东", "product": "A"}
        ck = _get_col_key(item, ["product"], "(空白)")
        self.assertEqual(ck, ("A",))

    def test_multiple_fields(self):
        item = {"region": "华东", "product": "A"}
        ck = _get_col_key(item, ["region", "product"], "(空白)")
        self.assertEqual(ck, ("华东", "A"))

    def test_empty_field_value(self):
        item = {"product": None}
        ck = _get_col_key(item, ["product"], "(空)")
        self.assertEqual(ck, ("(空)",))

    def test_no_col_fields(self):
        item = {"product": "A"}
        ck = _get_col_key(item, [], "(空白)")
        self.assertEqual(ck, ("",))


class TestBuildPivotTable(unittest.TestCase):
    def test_basic_pivot(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
        }
        result = build_pivot_table(SAMPLE_DATA, pivot_config, [])
        self.assertIn("row_keys", result)
        self.assertIn("col_keys", result)
        self.assertIn("data", result)
        self.assertEqual(len(result["row_keys"]), 3)
        self.assertEqual(len(result["col_keys"]), 2)

    def test_pivot_with_non_dict_items(self):
        data = SAMPLE_DATA + ["not_a_dict", None]
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
        }
        result = build_pivot_table(data, pivot_config, [])
        self.assertEqual(len(result["row_keys"]), 3)

    def test_pivot_no_row_fields(self):
        pivot_config = {
            "row_fields": [],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
        }
        result = build_pivot_table(SAMPLE_DATA, pivot_config, [])
        self.assertEqual(len(result["row_keys"]), 1)
        self.assertEqual(result["row_keys"][0], ())

    def test_pivot_no_col_fields(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": [],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
        }
        result = build_pivot_table(SAMPLE_DATA, pivot_config, [])
        self.assertEqual(len(result["col_keys"]), 1)
        self.assertEqual(result["col_keys"][0], ("",))

    def test_pivot_multiple_value_fields(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [
                {"field": "sales", "aggregate": "sum"},
                {"field": "quantity", "aggregate": "average"},
            ],
        }
        result = build_pivot_table(SAMPLE_DATA, pivot_config, [])
        self.assertIn("value_fields", result)
        self.assertEqual(len(result["value_fields"]), 2)

    def test_pivot_with_empty_label(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
            "empty_value_label": "N/A",
        }
        result = build_pivot_table(SAMPLE_DATA, pivot_config, [])
        self.assertEqual(result["empty_label"], "N/A")


class TestComputePivotValue(unittest.TestCase):
    def setUp(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
        }
        self.pivot_result = build_pivot_table(SAMPLE_DATA, pivot_config, [])
        self.vf = {"field": "sales", "aggregate": "sum"}

    def test_existing_value(self):
        rk = ("华东",)
        ck = ("A",)
        val = _compute_pivot_value(self.pivot_result, rk, ck, self.vf)
        self.assertEqual(val, 100)

    def test_nonexistent_row(self):
        rk = ("不存在",)
        ck = ("A",)
        val = _compute_pivot_value(self.pivot_result, rk, ck, self.vf)
        self.assertEqual(val, 0)

    def test_nonexistent_col(self):
        rk = ("华东",)
        ck = ("不存在",)
        val = _compute_pivot_value(self.pivot_result, rk, ck, self.vf)
        self.assertEqual(val, 0)


class TestComputeRowTotal(unittest.TestCase):
    def setUp(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [
                {"field": "sales", "aggregate": "sum"},
                {"field": "sales", "aggregate": "count"},
            ],
        }
        self.pivot_result = build_pivot_table(SAMPLE_DATA, pivot_config, [])

    def test_row_total_sum(self):
        vf = {"field": "sales", "aggregate": "sum"}
        rk = ("华东",)
        total = _compute_row_total(self.pivot_result, rk, vf)
        self.assertEqual(total, 300)

    def test_row_total_count(self):
        vf = {"field": "sales", "aggregate": "count"}
        rk = ("华东",)
        total = _compute_row_total(self.pivot_result, rk, vf)
        self.assertEqual(total, 2)


class TestComputeColTotal(unittest.TestCase):
    def setUp(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [
                {"field": "sales", "aggregate": "sum"},
                {"field": "sales", "aggregate": "count"},
            ],
        }
        self.pivot_result = build_pivot_table(SAMPLE_DATA, pivot_config, [])

    def test_col_total_sum(self):
        vf = {"field": "sales", "aggregate": "sum"}
        ck = ("A",)
        total = _compute_col_total(self.pivot_result, ck, vf)
        self.assertEqual(total, 550)

    def test_col_total_count(self):
        vf = {"field": "sales", "aggregate": "count"}
        ck = ("A",)
        total = _compute_col_total(self.pivot_result, ck, vf)
        self.assertEqual(total, 3)


class TestComputeGrandTotal(unittest.TestCase):
    def setUp(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [
                {"field": "sales", "aggregate": "sum"},
                {"field": "sales", "aggregate": "count"},
            ],
        }
        self.pivot_result = build_pivot_table(SAMPLE_DATA, pivot_config, [])

    def test_grand_total_sum(self):
        vf = {"field": "sales", "aggregate": "sum"}
        total = _compute_grand_total(self.pivot_result, vf)
        self.assertEqual(total, 1000)

    def test_grand_total_count(self):
        vf = {"field": "sales", "aggregate": "count"}
        total = _compute_grand_total(self.pivot_result, vf)
        self.assertEqual(total, 5)


class TestFormatValue(unittest.TestCase):
    def test_int_float(self):
        self.assertEqual(_format_value(10.0, "sum"), 10)

    def test_float_with_decimal(self):
        self.assertEqual(_format_value(10.567, "sum"), 10.57)

    def test_int_value(self):
        self.assertEqual(_format_value(10, "sum"), 10)

    def test_string_value(self):
        self.assertEqual(_format_value("abc", "sum"), "abc")


class TestCreatePivotSheet(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.headers = [
            {"key": "region", "label": "地区"},
            {"key": "product", "label": "产品"},
            {"key": "sales", "label": "销售额"},
        ]

    def test_create_basic_pivot_sheet(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
            "sheet_name": "数据透视表",
        }
        pivot_result = build_pivot_table(SAMPLE_DATA, pivot_config, self.headers)
        ws = create_pivot_sheet(self.wb, pivot_result, pivot_config, self.headers)
        self.assertIsNotNone(ws)
        self.assertEqual(ws.title, "数据透视表")

    def test_create_pivot_no_totals(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
            "show_row_totals": False,
            "show_column_totals": False,
        }
        pivot_result = build_pivot_table(SAMPLE_DATA, pivot_config, self.headers)
        ws = create_pivot_sheet(self.wb, pivot_result, pivot_config, self.headers)
        self.assertIsNotNone(ws)

    def test_create_pivot_no_style(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
            "apply_style": False,
        }
        pivot_result = build_pivot_table(SAMPLE_DATA, pivot_config, self.headers)
        ws = create_pivot_sheet(self.wb, pivot_result, pivot_config, self.headers)
        self.assertIsNotNone(ws)

    def test_create_pivot_single_value_field(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
        }
        pivot_result = build_pivot_table(SAMPLE_DATA, pivot_config, self.headers)
        ws = create_pivot_sheet(self.wb, pivot_result, pivot_config, self.headers)
        self.assertIsNotNone(ws)

    def test_create_pivot_multiple_value_fields(self):
        pivot_config = {
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [
                {"field": "sales", "aggregate": "sum"},
                {"field": "quantity", "aggregate": "average"},
            ],
        }
        headers = self.headers + [{"key": "quantity", "label": "数量"}]
        pivot_result = build_pivot_table(SAMPLE_DATA, pivot_config, headers)
        ws = create_pivot_sheet(self.wb, pivot_result, pivot_config, headers)
        self.assertIsNotNone(ws)


class TestAddPivotTableToWorkbook(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.config = get_default_config()
        self.headers = [
            {"key": "region", "label": "地区"},
            {"key": "product", "label": "产品"},
            {"key": "sales", "label": "销售额"},
        ]

    def test_add_pivot_table_enabled(self):
        self.config["pivot_config"] = {
            "enabled": True,
            "row_fields": ["region"],
            "column_fields": ["product"],
            "value_fields": [{"field": "sales", "aggregate": "sum"}],
        }
        result = add_pivot_table_to_workbook(self.wb, SAMPLE_DATA, self.headers, self.config)
        self.assertIsNotNone(result)
        self.assertIn("数据透视表", self.wb.sheetnames)

    def test_add_pivot_table_disabled(self):
        self.config["pivot_config"] = {"enabled": False}
        result = add_pivot_table_to_workbook(self.wb, SAMPLE_DATA, self.headers, self.config)
        self.assertIsNone(result)

    def test_add_pivot_table_no_config(self):
        result = add_pivot_table_to_workbook(self.wb, SAMPLE_DATA, self.headers, self.config)
        self.assertIsNone(result)


class TestApplyValidationMarks(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.headers = [{"key": "name", "label": "姓名"}, {"key": "age", "label": "年龄"}]
        self.ws.append(["姓名", "年龄"])
        for i in range(3):
            self.ws.append([f"用户{i}", 20 + i])
        self.original_indices = [0, 1, 2]

    def _make_result_with_errors(self, errors_data):
        result = ValidationResult()
        for row_idx, field_key, value, rule in errors_data:
            error = ValidationError(
                row_index=row_idx,
                rule=rule,
                value=value,
                field_key=field_key,
            )
            result.add_error(error)
        return result

    def test_with_marked_rows(self):
        rule = ValidationRule(
            field="age",
            rule_type="range",
            params={"min": 0, "max": 100},
            message="年龄范围错误",
            on_fail=ON_FAIL_MARK,
        )
        result = self._make_result_with_errors([
            (1, "age", 150, rule),
        ])
        _apply_validation_marks(self.ws, result, self.headers, self.original_indices)
        cell = self.ws.cell(row=3, column=1)
        self.assertIsNotNone(cell.comment)

    def test_no_marked_rows(self):
        result = ValidationResult()
        _apply_validation_marks(self.ws, result, self.headers, self.original_indices)
        cell = self.ws.cell(row=2, column=1)
        self.assertIsNone(cell.comment)

    def test_marked_row_not_in_original_indices(self):
        rule = ValidationRule(
            field="age",
            rule_type="range",
            params={"min": 0, "max": 100},
            message="年龄范围错误",
            on_fail=ON_FAIL_MARK,
        )
        result = self._make_result_with_errors([
            (5, "age", 150, rule),
        ])
        _apply_validation_marks(self.ws, result, self.headers, self.original_indices)
        cell = self.ws.cell(row=2, column=1)
        self.assertIsNone(cell.comment)

    def test_long_comment_truncated(self):
        rule = ValidationRule(
            field="age",
            rule_type="range",
            params={"min": 0, "max": 100},
            message="年龄范围错误" * 50,
            on_fail=ON_FAIL_MARK,
        )
        result = self._make_result_with_errors([
            (1, "age", 150, rule),
        ])
        _apply_validation_marks(self.ws, result, self.headers, self.original_indices)
        cell = self.ws.cell(row=3, column=1)
        self.assertTrue(len(cell.comment.text) <= 203)


class TestPrintValidationErrors(unittest.TestCase):
    def _make_result_with_errors(self, errors_data):
        result = ValidationResult()
        for row_idx, field_key, value, rule in errors_data:
            error = ValidationError(
                row_index=row_idx,
                rule=rule,
                value=value,
                field_key=field_key,
            )
            result.add_error(error)
        return result

    def test_no_errors(self):
        result = ValidationResult()
        _print_validation_errors(result, [])

    def test_with_errors(self):
        rule = ValidationRule(
            field="age",
            rule_type="range",
            params={"min": 0, "max": 100},
            message="年龄范围错误",
        )
        result = self._make_result_with_errors([
            (0, "age", 150, rule),
            (1, "name", None, rule),
        ])
        _print_validation_errors(result, [{"age": 150}, {"name": None}])

    def test_errors_exceed_max_display(self):
        rule = ValidationRule(
            field="age",
            rule_type="range",
            params={"min": 0, "max": 100},
            message="年龄范围错误",
        )
        errors_data = [(i, "age", 100 + i, rule) for i in range(15)]
        result = self._make_result_with_errors(errors_data)
        _print_validation_errors(result, [{"age": 100 + i} for i in range(15)], max_display=10)


class TestMatchCustomRule(unittest.TestCase):
    def test_values_list_match(self):
        rule = {"values": ["test", "hello"]}
        self.assertTrue(_match_custom_rule("test", rule))

    def test_values_list_no_match(self):
        rule = {"values": ["test", "hello"]}
        self.assertFalse(_match_custom_rule("other", rule))

    def test_values_single_match(self):
        rule = {"values": "test"}
        self.assertTrue(_match_custom_rule("test", rule))

    def test_condition_match(self):
        rule = {"condition": "value > 10"}
        self.assertTrue(_match_custom_rule(15, rule))

    def test_condition_no_match(self):
        rule = {"condition": "value > 10"}
        self.assertFalse(_match_custom_rule(5, rule))

    def test_condition_exception(self):
        rule = {"condition": "value + undefined_var"}
        self.assertFalse(_match_custom_rule(10, rule))

    def test_min_max_in_range(self):
        rule = {"min": 10, "max": 20}
        self.assertTrue(_match_custom_rule(15, rule))

    def test_min_max_below_min(self):
        rule = {"min": 10, "max": 20}
        self.assertFalse(_match_custom_rule(5, rule))

    def test_min_max_above_max(self):
        rule = {"min": 10, "max": 20}
        self.assertFalse(_match_custom_rule(25, rule))

    def test_min_only(self):
        rule = {"min": 10}
        self.assertTrue(_match_custom_rule(15, rule))
        self.assertFalse(_match_custom_rule(5, rule))

    def test_max_only(self):
        rule = {"max": 20}
        self.assertTrue(_match_custom_rule(15, rule))
        self.assertFalse(_match_custom_rule(25, rule))

    def test_include_min_false(self):
        rule = {"min": 10, "include_min": False}
        self.assertFalse(_match_custom_rule(10, rule))
        self.assertTrue(_match_custom_rule(11, rule))

    def test_include_max_true(self):
        rule = {"max": 20, "include_max": True}
        self.assertTrue(_match_custom_rule(20, rule))

    def test_none_value_with_min_max(self):
        rule = {"min": 0, "max": 20}
        self.assertTrue(_match_custom_rule(None, rule))

    def test_empty_string_value(self):
        rule = {"min": 0, "max": 20}
        self.assertTrue(_match_custom_rule("", rule))

    def test_invalid_numeric_value(self):
        rule = {"min": 10, "max": 20}
        self.assertFalse(_match_custom_rule("abc", rule))

    def test_no_recognized_keys(self):
        rule = {"unknown_key": "value"}
        self.assertFalse(_match_custom_rule("test", rule))


class TestMatchRangeGroup(unittest.TestCase):
    def test_match_in_range(self):
        group = {"min": 10, "max": 20, "label": "10-20"}
        self.assertTrue(_match_range_group(15, group))

    def test_match_at_min(self):
        group = {"min": 10, "max": 20, "label": "10-20"}
        self.assertTrue(_match_range_group(10, group))

    def test_match_at_max_default_include_max_false(self):
        group = {"min": 10, "max": 20, "label": "10-20"}
        self.assertFalse(_match_range_group(20, group))

    def test_match_at_max_include_max_true(self):
        group = {"min": 10, "max": 20, "label": "10-20", "include_max": True}
        self.assertTrue(_match_range_group(20, group))

    def test_below_range(self):
        group = {"min": 10, "max": 20, "label": "10-20"}
        self.assertFalse(_match_range_group(5, group))

    def test_above_range(self):
        group = {"min": 10, "max": 20, "label": "10-20"}
        self.assertFalse(_match_range_group(25, group))

    def test_none_value(self):
        group = {"min": 10, "max": 20, "label": "10-20"}
        self.assertFalse(_match_range_group(None, group))

    def test_empty_string_value(self):
        group = {"min": 10, "max": 20, "label": "10-20"}
        self.assertFalse(_match_range_group("", group))

    def test_none_min(self):
        group = {"min": None, "max": 20, "label": "<=20"}
        self.assertTrue(_match_range_group(10, group))

    def test_none_max(self):
        group = {"min": 10, "max": None, "label": ">=10"}
        self.assertTrue(_match_range_group(20, group))

    def test_invalid_value_type(self):
        group = {"min": 10, "max": 20, "label": "10-20"}
        self.assertFalse(_match_range_group("abc", group))

    def test_include_min_false(self):
        group = {"min": 10, "max": 20, "label": "10-20", "include_min": False}
        self.assertFalse(_match_range_group(10, group))
        self.assertTrue(_match_range_group(11, group))


class TestSanitizeSheetName(unittest.TestCase):
    def test_normal_name(self):
        self.assertEqual(_sanitize_sheet_name("Sheet1"), "Sheet1")

    def test_name_too_long(self):
        long_name = "A" * 50
        result = _sanitize_sheet_name(long_name)
        self.assertEqual(len(result), 31)

    def test_invalid_chars(self):
        self.assertEqual(_sanitize_sheet_name("Sheet/1"), "Sheet_1")

    def test_contains_colon(self):
        self.assertEqual(_sanitize_sheet_name("Sheet:1"), "Sheet_1")

    def test_contains_asterisk(self):
        self.assertEqual(_sanitize_sheet_name("Sheet*1"), "Sheet_1")

    def test_contains_question(self):
        self.assertEqual(_sanitize_sheet_name("Sheet?1"), "Sheet_1")

    def test_contains_brackets(self):
        self.assertEqual(_sanitize_sheet_name("Sheet[1]"), "Sheet_1_")


class TestRenderSheetName(unittest.TestCase):
    def test_with_value(self):
        name = _render_sheet_name("{value}", "华东", 0, 3)
        self.assertEqual(name, "华东")

    def test_with_index(self):
        name = _render_sheet_name("Sheet{index}", "", 0, 3)
        self.assertEqual(name, "Sheet0")

    def test_with_value_and_index(self):
        name = _render_sheet_name("{value}_{index}", "华东", 0, 3)
        self.assertEqual(name, "华东_0")

    def test_with_count(self):
        name = _render_sheet_name("共{count}个", "华东", 0, 3)
        self.assertEqual(name, "共3个")

    def test_with_num(self):
        name = _render_sheet_name("第{num}个", "华东", 1, 3)
        self.assertEqual(name, "第1个")

    def test_empty_value(self):
        name = _render_sheet_name("{value}", "", 0, 3, empty_label="未分类")
        self.assertEqual(name, "未分类")

    def test_none_value(self):
        name = _render_sheet_name("{value}", None, 0, 3, empty_label="未分类")
        self.assertEqual(name, "未分类")

    def test_invalid_template_key(self):
        name = _render_sheet_name("{unknown}", "华东", 0, 3)
        self.assertEqual(name, "华东")


class TestSplitDataByField(unittest.TestCase):
    def test_split_by_value(self):
        split_config = {"split_rule": "by_value"}
        result = split_data_by_field(SAMPLE_DATA, "region", split_config)
        self.assertEqual(len(result), 3)
        self.assertIn("华东", [g["name"] for g in result])

    def test_split_by_range(self):
        split_config = {
            "split_rule": "by_range",
            "range_groups": [
                {"min": None, "max": 150, "name": "低", "include_max": True},
                {"min": 150, "max": 300, "name": "中"},
                {"min": 300, "max": None, "name": "高"},
            ],
        }
        result = split_data_by_field(SAMPLE_DATA, "sales", split_config)
        self.assertEqual(len(result), 3)

    def test_split_by_custom(self):
        split_config = {
            "split_rule": "by_custom",
            "custom_rules": [
                {"values": ["华东", "华北"], "name": "北方"},
                {"values": ["华南"], "name": "南方"},
            ],
        }
        result = split_data_by_field(SAMPLE_DATA, "region", split_config)
        self.assertEqual(len(result), 2)

    def test_split_with_default(self):
        split_config = {
            "split_rule": "by_custom",
            "custom_rules": [
                {"values": ["华东"], "name": "华东"},
            ],
            "fallback_group_name": "其他",
        }
        result = split_data_by_field(SAMPLE_DATA, "region", split_config)
        self.assertEqual(len(result), 2)
        group_names = [g["name"] for g in result]
        self.assertIn("其他", group_names)

    def test_split_non_dict_items(self):
        data = SAMPLE_DATA + ["not_dict", None]
        split_config = {"split_rule": "by_value"}
        result = split_data_by_field(data, "region", split_config)
        self.assertGreater(len(result), 0)

    def test_split_invalid_type(self):
        split_config = {"split_rule": "unknown_type"}
        result = split_data_by_field(SAMPLE_DATA, "region", split_config)
        self.assertEqual(len(result), 0)

    def test_split_with_empty_value_label(self):
        data = [{"region": ""}, {"region": None}, {"region": "华东"}]
        split_config = {"split_rule": "by_value", "empty_value_label": "空值"}
        result = split_data_by_field(data, "region", split_config)
        group_names = [g["name"] for g in result]
        self.assertIn("空值", group_names)


if __name__ == "__main__":
    unittest.main()
